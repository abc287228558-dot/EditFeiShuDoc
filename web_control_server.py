import argparse
import csv
import io
import json
import os
import re
import secrets
import subprocess
import sys
import threading
import time
from datetime import datetime
from http.server import BaseHTTPRequestHandler, HTTPServer
from typing import Any, Dict, List, Optional, Tuple
from urllib.parse import parse_qs, quote, urlparse

import openpyxl
from openpyxl.workbook.workbook import Workbook


SERVER_VERSION = time.strftime("%Y%m%d-%H%M%S", time.localtime())


def _now_iso() -> str:
    return datetime.now().strftime("%Y-%m-%d %H:%M:%S")


def load_json(path: str) -> Dict[str, Any]:
    with open(path, "r", encoding="utf-8") as f:
        return json.load(f)


def save_json(path: str, data: Dict[str, Any]) -> None:
    with open(path, "w", encoding="utf-8") as f:
        json.dump(data, f, ensure_ascii=False, indent=2)


def _resolve_path(base_file: str, maybe_relative: str) -> str:
    if not maybe_relative:
        return ""
    if os.path.isabs(maybe_relative):
        return maybe_relative
    base_dir = os.path.dirname(os.path.abspath(base_file))
    if os.path.basename(base_dir) == ".state" and maybe_relative.startswith(f".state{os.sep}"):
        base_dir = os.path.dirname(base_dir)
    return os.path.join(base_dir, maybe_relative)


def load_accounts_from_anchor_csv(path: str) -> List[str]:
    if not path or not os.path.exists(path):
        return []
    with open(path, "r", encoding="utf-8-sig") as f:
        reader = csv.DictReader(f)
        if reader.fieldnames:
            reader.fieldnames = [n.strip() for n in reader.fieldnames]
        accounts: List[str] = []
        for r in reader:
            acct = ((r.get("直播账号") or r.get("\ufeff直播账号") or "")).strip()
            if acct:
                accounts.append(acct)
        return accounts


def _read_csv_rows(path: str) -> List[List[str]]:
    if not path or not os.path.exists(path):
        return []
    rows: List[List[str]] = []
    header = ["直播账号", "快手ID", "手机号码", "密码", "主播"]
    with open(path, "r", encoding="utf-8-sig", newline="") as f:
        reader = csv.reader(f)
        for i, r in enumerate(reader):
            if not isinstance(r, list):
                continue
            row = [str(x) if x is not None else "" for x in r]
            if i == 0 and row[: len(header)] == header:
                continue
            rows.append(row)
    return rows


def _anchor_map_csv_path(cfg: Dict[str, Any], args: Any) -> str:
    anchor_map_csv_raw = (cfg.get("mapping") or {}).get("anchor_map_csv") or ""
    return _resolve_path(args.config, str(anchor_map_csv_raw))


def _wenzong_anchor_map_csv_path(cfg: Dict[str, Any], args: Any) -> str:
    raw = (cfg.get("mapping") or {}).get("wenzong_anchor_map_csv") or ""
    return _resolve_path(args.config, str(raw))


def _inactive_csv_path(cfg: Dict[str, Any], args: Any) -> str:
    inactive_csv_raw = (cfg.get("mapping") or {}).get("inactive_csv") or ""
    return _resolve_path(args.config, str(inactive_csv_raw))


def _json_response(handler: BaseHTTPRequestHandler, status: int, payload: Dict[str, Any]) -> None:
    body = json.dumps(payload, ensure_ascii=False).encode("utf-8")
    handler.send_response(status)
    handler.send_header("Content-Type", "application/json; charset=utf-8")
    handler.send_header("Cache-Control", "no-store")
    handler.send_header("Pragma", "no-cache")
    handler.send_header("Expires", "0")
    handler.send_header("Content-Length", str(len(body)))
    handler.end_headers()
    handler.wfile.write(body)


def _xlsx_bytes_from_rows(rows: List[List[str]]) -> bytes:
    wb: Workbook = openpyxl.Workbook()
    ws = wb.active
    ws.title = "导入"
    ws.append(["姓名", "账号", "别名", "部门", "手机"])
    for r in rows or []:
        if not isinstance(r, list):
            continue
        r5 = [(str(r[j]).strip() if j < len(r) and r[j] is not None else "") for j in range(5)]
        if any(r5):
            ws.append(r5)
    bio = io.BytesIO()
    wb.save(bio)
    return bio.getvalue()


def _find_header_row(values: List[Tuple[Any, ...]], headers: List[str], max_scan: int = 50) -> Tuple[int, Dict[str, int]]:
    scan = values[: max_scan if max_scan > 0 else len(values)]
    for ridx, row in enumerate(scan):
        cells = [str(c).strip() if c is not None else "" for c in (row or [])]
        idx_map: Dict[str, int] = {}
        ok = True
        for h in headers:
            try:
                idx_map[h] = cells.index(h)
            except ValueError:
                ok = False
                break
        if ok:
            return ridx, idx_map
    return -1, {}


def _parse_xlsx_rows(data: bytes) -> List[List[str]]:
    wb = openpyxl.load_workbook(io.BytesIO(data), data_only=True)
    ws = wb.active
    values = list(ws.iter_rows(values_only=True))
    if not values:
        return []

    required = ["姓名", "账号", "别名", "手机"]
    header_row, idx_req = _find_header_row(values, required, max_scan=80)
    if header_row < 0:
        # Fallback: treat as plain 5-column table from first row.
        out: List[List[str]] = []
        for row in values:
            if row is None:
                continue
            r5 = [("" if j >= len(row) or row[j] is None else str(row[j]).strip()) for j in range(5)]
            if any(r5):
                out.append(r5)
        return out

    # Optional dept column
    header_cells = [str(c).strip() if c is not None else "" for c in (values[header_row] or [])]
    dept_idx = -1
    try:
        dept_idx = header_cells.index("部门")
    except ValueError:
        dept_idx = -1

    # Build output in fixed 5-col order: 姓名/账号/别名/部门/手机
    out: List[List[str]] = []
    for row in values[header_row + 1 :]:
        if row is None:
            continue
        name = row[idx_req["姓名"]] if idx_req["姓名"] < len(row) else ""
        acct = row[idx_req["账号"]] if idx_req["账号"] < len(row) else ""
        alias = row[idx_req["别名"]] if idx_req["别名"] < len(row) else ""
        phone = row[idx_req["手机"]] if idx_req["手机"] < len(row) else ""
        dept = row[dept_idx] if (dept_idx >= 0 and dept_idx < len(row)) else ""
        r5 = [
            "" if name is None else str(name).strip(),
            "" if acct is None else str(acct).strip(),
            "" if alias is None else str(alias).strip(),
            "" if dept is None else str(dept).strip(),
            "" if phone is None else str(phone).strip(),
        ]
        if any(r5):
            out.append(r5)
    return out


def _user_contact_template_xlsx_path(cfg: Dict[str, Any], args: Any) -> str:
    raw = ((cfg.get("mapping") or {}).get("user_contact_template_xlsx") or "2.28琳琳.xlsx")
    raw_s = str(raw)
    # First try resolving relative to the config file path.
    p1 = _resolve_path(args.config, raw_s)
    if p1 and os.path.exists(p1):
        return p1
    # Fallback: resolve relative to current working directory.
    if raw_s and not os.path.isabs(raw_s):
        p2 = os.path.abspath(raw_s)
        if os.path.exists(p2):
            return p2
    # Fallback: resolve relative to this script directory.
    if raw_s and not os.path.isabs(raw_s):
        p3 = os.path.join(os.path.dirname(os.path.abspath(__file__)), raw_s)
        if os.path.exists(p3):
            return p3
    return p1


def _xlsx_bytes_from_rows_template(cfg: Dict[str, Any], args: Any, rows: List[List[str]]) -> bytes:
    template_path = _user_contact_template_xlsx_path(cfg, args)
    if not template_path or not os.path.exists(template_path):
        raise FileNotFoundError(f"template_not_found: {template_path}")

    if not rows:
        with open(template_path, "rb") as f:
            return f.read()

    wb = openpyxl.load_workbook(template_path)
    ws = wb.active
    values = list(ws.iter_rows(values_only=True))
    header_row, idx_req = _find_header_row(values, ["姓名", "账号", "别名", "手机"], max_scan=120)
    if header_row < 0:
        raise ValueError("template_header_not_found")

    header_cells = [str(c).strip() if c is not None else "" for c in (values[header_row] or [])]
    dept_idx = header_cells.index("部门") if "部门" in header_cells else -1

    # Convert 0-based indices to 1-based excel columns
    col_name = idx_req["姓名"] + 1
    col_acct = idx_req["账号"] + 1
    col_alias = idx_req["别名"] + 1
    col_phone = idx_req["手机"] + 1
    col_dept = (dept_idx + 1) if dept_idx >= 0 else None

    start_row = header_row + 2

    # Clear only the rows that will be overwritten (existing data range + new rows)
    cols = [col_name, col_acct, col_alias, col_phone] + ([col_dept] if col_dept else [])
    existing_end = start_row - 1
    for r in range(ws.max_row, start_row - 1, -1):
        for c in cols:
            v = ws.cell(r, c).value
            if v is not None and str(v).strip() != "":
                existing_end = r
                break
        if existing_end >= start_row:
            break

    new_end = start_row + max(len(rows) - 1, 0)
    clear_end = max(existing_end, new_end)
    for r in range(start_row, clear_end + 1):
        for c in cols:
            ws.cell(r, c).value = None

    for i, r in enumerate(rows or []):
        if not isinstance(r, list):
            continue
        r5 = [(str(r[j]).strip() if j < len(r) and r[j] is not None else "") for j in range(5)]
        if not any(r5):
            continue
        rr = start_row + i
        ws.cell(rr, col_name).value = r5[0]
        ws.cell(rr, col_acct).value = r5[1]
        ws.cell(rr, col_alias).value = r5[2]
        if col_dept:
            ws.cell(rr, col_dept).value = r5[3]
        ws.cell(rr, col_phone).value = r5[4]

    bio = io.BytesIO()
    wb.save(bio)
    return bio.getvalue()


def _apply_dept(rows: List[List[str]], dept: str) -> List[List[str]]:
    d = (dept or "").strip()
    out: List[List[str]] = []
    for r in rows or []:
        if not isinstance(r, list):
            continue
        r5 = [(str(r[j]).strip() if j < len(r) and r[j] is not None else "") for j in range(5)]
        r5[3] = d
        if any(r5):
            out.append(r5)
    return out


def _bytes_response(
    handler: BaseHTTPRequestHandler,
    status: int,
    body: bytes,
    *,
    content_type: str,
    filename: str = "",
) -> None:
    handler.send_response(status)
    handler.send_header("Content-Type", content_type)
    handler.send_header("Cache-Control", "no-store")
    handler.send_header("Pragma", "no-cache")
    handler.send_header("Expires", "0")
    if filename:
        safe = filename.replace('"', "")
        ascii_fallback = re.sub(r"[^A-Za-z0-9._-]+", "_", safe).strip("._-") or "download"
        quoted = quote(safe, safe="")
        handler.send_header(
            "Content-Disposition",
            f"attachment; filename=\"{ascii_fallback}\"; filename*=UTF-8''{quoted}",
        )
    handler.send_header("Content-Length", str(len(body)))
    handler.end_headers()
    handler.wfile.write(body)


def _html_response(handler: BaseHTTPRequestHandler, status: int, html: str) -> None:
    raw = html.encode("utf-8")
    handler.send_response(status)
    handler.send_header("Content-Type", "text/html; charset=utf-8")
    handler.send_header("Cache-Control", "no-store")
    handler.send_header("Pragma", "no-cache")
    handler.send_header("Expires", "0")
    handler.send_header("Content-Length", str(len(raw)))
    handler.end_headers()
    handler.wfile.write(raw)


def _text_response(handler: BaseHTTPRequestHandler, status: int, text: str) -> None:
    raw = text.encode("utf-8")
    handler.send_response(status)
    handler.send_header("Content-Type", "text/plain; charset=utf-8")
    handler.send_header("Cache-Control", "no-store")
    handler.send_header("Pragma", "no-cache")
    handler.send_header("Expires", "0")
    handler.send_header("Content-Length", str(len(raw)))
    handler.end_headers()
    handler.wfile.write(raw)


def _redirect(handler: BaseHTTPRequestHandler, location: str) -> None:
    handler.send_response(302)
    handler.send_header("Location", location)
    handler.end_headers()


def _html_escape(s: str) -> str:
    return (
        (s or "")
        .replace("&", "&amp;")
        .replace("<", "&lt;")
        .replace(">", "&gt;")
        .replace('"', "&quot;")
        .replace("'", "&#39;")
    )


def _cleanup_keep_latest_files(dir_path: str, *, keep: int, exts: Optional[List[str]] = None) -> None:
    if keep <= 0:
        return

    if not dir_path or not os.path.isdir(dir_path):
        return
    try:
        items = []
        for name in os.listdir(dir_path):
            p = os.path.join(dir_path, name)
            if not os.path.isfile(p):
                continue
            if exts is not None:
                _, ext = os.path.splitext(name)
                if ext.lower() not in {e.lower() for e in exts}:
                    continue
            try:
                st = os.stat(p)
            except Exception:
                continue
            items.append((st.st_mtime, p))

        items.sort(reverse=True)
        for _, p in items[keep:]:
            try:
                os.remove(p)
            except Exception:
                pass
    except Exception:
        return


def _clean_browser_profile_caches(profile_dirs: List[str]) -> int:
    try:
        import shutil
        from pathlib import Path

        cache_paths = [
            'Default/Cache',
            'Default/Code Cache',
            'Default/GPUCache',
            'Default/DawnWebGPUCache',
            'Default/DawnGraphiteCache',
            'Default/Service Worker/CacheStorage',
            'Default/Service Worker/ScriptCache',
            'Default/Service Worker/Database',
            'Default/Shared Dictionary/cache',
            'Default/optimization_guide_hint_cache_store',
            'Default/Favicons',
            'Default/Favicons-journal',
            'Default/Favicons-wal',
            'Default/Favicons-shm',
            'Default/History',
            'Default/History-journal',
            'Default/History-wal',
            'Default/History-shm',
            'BrowserMetrics',
            'GrShaderCache',
            'ShaderCache',
            'GraphiteDawnCache',
            'Cache',
            'Code Cache',
            'GPUCache',
            'segmentation_platform',
        ]

        total_freed = 0
        for profile_dir in profile_dirs:
            profile_path = Path(profile_dir)
            if not profile_path.exists() or (not profile_path.is_dir()):
                continue

            # Support two layouts:
            # 1) profile_path contains many profile subdirs (e.g. kuaishou_profiles/<ks_id>)
            # 2) profile_path itself is a user-data-dir containing Default/... (e.g. feishu_profile)
            candidate_roots: List[Path] = []
            try:
                if (profile_path / 'Default').exists():
                    candidate_roots.append(profile_path)
                else:
                    # If any cache path exists directly under profile_path, treat it as a root too.
                    for cache_rel_path in cache_paths:
                        if (profile_path / cache_rel_path).exists():
                            candidate_roots.append(profile_path)
                            break
            except Exception:
                pass

            try:
                for item in profile_path.iterdir():
                    if item.is_dir():
                        candidate_roots.append(item)
            except Exception:
                pass

            seen_roots = set()
            roots2: List[Path] = []
            for r in candidate_roots:
                try:
                    key = str(r.resolve())
                except Exception:
                    key = str(r)
                if key in seen_roots:
                    continue
                seen_roots.add(key)
                roots2.append(r)

            for root in roots2:
                for cache_rel_path in cache_paths:
                    cache_full_path = root / cache_rel_path
                    if not cache_full_path.exists():
                        continue
                    try:
                        if cache_full_path.is_file():
                            size = cache_full_path.stat().st_size
                            try:
                                cache_full_path.unlink()
                            except Exception:
                                try:
                                    cache_full_path.unlink(missing_ok=True)
                                except Exception:
                                    pass
                            total_freed += size
                        else:
                            size = sum(
                                f.stat().st_size
                                for f in cache_full_path.rglob('*')
                                if f.is_file()
                            )
                            shutil.rmtree(cache_full_path, ignore_errors=True)
                            total_freed += size
                    except Exception:
                        pass
        return int(total_freed)
    except Exception:
        return 0


def _now_ts() -> str:
    return datetime.now().strftime("%Y%m%d_%H%M%S")


def _chinese_financial_num(n: int) -> str:
    digits = ["零", "壹", "贰", "叁", "肆", "伍", "陆", "柒", "捌", "玖"]
    if n <= 0:
        return "零"
    if n < 10:
        return digits[n]
    if n < 20:
        if n == 10:
            return "拾"
        return "拾" + digits[n % 10]
    if n < 100:
        tens = n // 10
        ones = n % 10
        if ones == 0:
            return digits[tens] + "拾"
        return digits[tens] + "拾" + digits[ones]
    return f"第{n}"


def _chinese_simple_num(n: int) -> str:
    digits = ["零", "一", "二", "三", "四", "五", "六", "七", "八", "九"]
    if n <= 0:
        return "零"
    if n < 10:
        return digits[n]
    if n < 20:
        if n == 10:
            return "十"
        return "十" + digits[n % 10]
    if n < 100:
        tens = n // 10
        ones = n % 10
        if ones == 0:
            return digits[tens] + "十"
        return digits[tens] + "十" + digits[ones]
    return f"第{n}"


def _read_multipart_file(raw: bytes, content_type: str) -> Tuple[str, bytes]:
    if "multipart/form-data" not in (content_type or ""):
        raise ValueError("not_multipart")
    m = re.search(r"boundary=(.+)", content_type)
    if not m:
        raise ValueError("missing_boundary")
    boundary = m.group(1)
    boundary = boundary.strip().strip('"')
    if not boundary:
        raise ValueError("empty_boundary")
    marker = ("--" + boundary).encode("utf-8")
    parts = raw.split(marker)
    for p in parts:
        p = p.strip(b"\r\n")
        if not p or p == b"--":
            continue
        head, sep, body = p.partition(b"\r\n\r\n")
        if not sep:
            continue
        header_text = head.decode("utf-8", errors="ignore")
        if "Content-Disposition" not in header_text:
            continue
        if "filename=" not in header_text:
            continue
        m2 = re.search(r"filename=\"([^\"]*)\"", header_text)
        filename = m2.group(1) if m2 else ""
        body = body.rstrip(b"\r\n")
        if body.endswith(b"--"):
            body = body[:-2]
        if body:
            return filename, body
    raise ValueError("file_not_found")


def _load_user_contact_rows(csv_path: str) -> List[List[str]]:
    if not csv_path or not os.path.exists(csv_path):
        return []
    rows: List[List[str]] = []
    with open(csv_path, "r", encoding="utf-8-sig") as f:
        reader = csv.reader(f)
        for i, row in enumerate(reader):
            if i == 0:
                continue
            if not isinstance(row, list):
                continue
            row5 = [(row[j].strip() if j < len(row) and row[j] is not None else "") for j in range(5)]
            if any(row5):
                rows.append(row5)
    return rows


def _save_user_contact_rows(csv_path: str, rows: List[List[str]]) -> None:
    os.makedirs(os.path.dirname(csv_path) or ".", exist_ok=True)
    header = ["姓名", "账号", "别名", "部门", "手机"]
    with open(csv_path, "w", encoding="utf-8-sig", newline="") as f:
        w = csv.writer(f)
        w.writerow(header)
        for r in rows:
            if not isinstance(r, list):
                continue
            r5 = [(str(r[j]).strip() if j < len(r) and r[j] is not None else "") for j in range(5)]
            if any(r5):
                w.writerow(r5)


def _user_contact_csv_path(cfg: Dict[str, Any], args: Any) -> str:
    raw = ((cfg.get("mapping") or {}).get("user_contact_csv") or ".state/user_contact_import.csv")
    return _resolve_path(args.config, str(raw))


def _user_contact_seq_path(cfg: Dict[str, Any], args: Any) -> str:
    raw = ((cfg.get("mapping") or {}).get("user_contact_seq_json") or ".state/user_contact_seq.json")
    return _resolve_path(args.config, str(raw))


def _wenzong_user_contact_seq_path(cfg: Dict[str, Any], args: Any) -> str:
    raw = ((cfg.get("mapping") or {}).get("wenzong_user_contact_seq_json") or ".state/user_contact_seq.wenzong.json")
    return _resolve_path(args.config, str(raw))


def _user_contact_cfg_path(cfg: Dict[str, Any], args: Any) -> str:
    raw = ((cfg.get("mapping") or {}).get("user_contact_cfg_json") or ".state/user_contact_cfg.json")
    return _resolve_path(args.config, str(raw))


def _wenzong_user_contact_cfg_path(cfg: Dict[str, Any], args: Any) -> str:
    raw = ((cfg.get("mapping") or {}).get("wenzong_user_contact_cfg_json") or ".state/user_contact_cfg.wenzong.json")
    return _resolve_path(args.config, str(raw))


def _user_contact_export_dir(cfg: Dict[str, Any], args: Any) -> str:
    raw = ((cfg.get("mapping") or {}).get("user_contact_export_dir") or "")
    if not raw:
        return ""
    v = str(raw).strip()
    if not v:
        return ""
    return _resolve_path(args.config, v)


def _wenzong_user_contact_export_dir(cfg: Dict[str, Any], args: Any) -> str:
    raw = ((cfg.get("mapping") or {}).get("wenzong_user_contact_export_dir") or "")
    if not raw:
        return ""
    v = str(raw).strip()
    if not v:
        return ""
    return _resolve_path(args.config, v)


def _niu_table_json_path(cfg: Dict[str, Any], args: Any) -> str:
    raw = ((cfg.get("mapping") or {}).get("niu_table_json") or ".state/niu_table.json")
    return _resolve_path(args.config, str(raw))


def _load_niu_table(cfg: Dict[str, Any], args: Any) -> List[Dict[str, str]]:
    p = _niu_table_json_path(cfg, args)
    try:
        if not os.path.exists(p):
            return []
        with open(p, "r", encoding="utf-8") as f:
            data = json.load(f)
        if not isinstance(data, list):
            return []
        out: List[Dict[str, str]] = []
        for it in data:
            if not isinstance(it, dict):
                continue
            url = str(it.get("url") or "").strip()
            if not url:
                continue
            out.append({"url": url})
        return out
    except Exception:
        return []


def _save_niu_table(cfg: Dict[str, Any], args: Any, rows: List[Dict[str, str]]) -> None:
    p = _niu_table_json_path(cfg, args)
    os.makedirs(os.path.dirname(p) or ".", exist_ok=True)
    out: List[Dict[str, str]] = []
    for it in rows or []:
        if not isinstance(it, dict):
            continue
        url = str(it.get("url") or "").strip()
        if not url:
            continue
        out.append({"url": url})
    with open(p, "w", encoding="utf-8") as f:
        json.dump(out, f, ensure_ascii=False, indent=2)
        f.write("\n")


def _niu_bg_status_json_path(cfg: Dict[str, Any], args: Any) -> str:
    raw = ((cfg.get("mapping") or {}).get("niu_bg_status_json") or ".state/niu_bg_status.json")
    return _resolve_path(args.config, str(raw))


def _load_niu_bg_status(cfg: Dict[str, Any], args: Any) -> Dict[str, Any]:
    p = _niu_bg_status_json_path(cfg, args)
    try:
        if not os.path.exists(p):
            return {}
        with open(p, "r", encoding="utf-8") as f:
            data = json.load(f)
        return data if isinstance(data, dict) else {}
    except Exception:
        return {}


def _niu_runtime_status_json_path(args: Any) -> str:
    return _resolve_path(args.config, ".state/niu_runtime_status.json")


def _load_niu_runtime_status(args: Any) -> Dict[str, Any]:
    p = _niu_runtime_status_json_path(args)
    try:
        if not os.path.exists(p):
            return {}
        with open(p, "r", encoding="utf-8") as f:
            data = json.load(f)
        return data if isinstance(data, dict) else {}
    except Exception:
        return {}


def _kuaishou_runtime_status_json_path(args: Any) -> str:
    return _resolve_path(args.config, ".state/kuaishou_runtime_status.json")


def _load_kuaishou_runtime_status(args: Any) -> Dict[str, Any]:
    p = _kuaishou_runtime_status_json_path(args)
    try:
        if not os.path.exists(p):
            return {}
        with open(p, "r", encoding="utf-8") as f:
            data = json.load(f)
        return data if isinstance(data, dict) else {}
    except Exception:
        return {}


def _feishu_runtime_status_json_path(args: Any) -> str:
    return _resolve_path(args.config, ".state/feishu_runtime_status.json")


def _load_feishu_runtime_status(args: Any) -> Dict[str, Any]:
    p = _feishu_runtime_status_json_path(args)
    try:
        if not os.path.exists(p):
            return {}
        with open(p, "r", encoding="utf-8") as f:
            data = json.load(f)
        return data if isinstance(data, dict) else {}
    except Exception:
        return {}


def _niu_front_profile_dir(idx: int) -> str:
    return os.path.join(".state", "niu_pw_profiles_front", f"slot_{idx}")


def _niu_bg_profile_dir(idx: int) -> str:
    return os.path.join(".state", "niu_pw_profiles_bg", f"slot_{idx}")


def _kuaishou_front_profile_dir(ks_id: str) -> str:
    safe = re.sub(r"[^0-9A-Za-z._-]+", "_", str(ks_id or "").strip()) or "unknown"
    return os.path.join(".state", "kuaishou_profiles_front", safe)


def _kuaishou_bg_profile_dir(ks_id: str) -> str:
    safe = re.sub(r"[^0-9A-Za-z._-]+", "_", str(ks_id or "").strip()) or "unknown"
    return os.path.join(".state", "kuaishou_profiles_bg", safe)


def _kuaishou_bg_status_json_path(cfg: Dict[str, Any], args: Any) -> str:
    raw = ((cfg.get("mapping") or {}).get("kuaishou_bg_status_json") or ".state/kuaishou_bg_status.json")
    return _resolve_path(args.config, str(raw))


def _load_kuaishou_bg_status(cfg: Dict[str, Any], args: Any) -> Dict[str, Any]:
    p = _kuaishou_bg_status_json_path(cfg, args)
    try:
        if not os.path.exists(p):
            return {}
        with open(p, "r", encoding="utf-8") as f:
            data = json.load(f)
        return data if isinstance(data, dict) else {}
    except Exception:
        return {}


def _save_kuaishou_bg_status(cfg: Dict[str, Any], args: Any, data: Dict[str, Any]) -> None:
    p = _kuaishou_bg_status_json_path(cfg, args)
    os.makedirs(os.path.dirname(p) or ".", exist_ok=True)
    with open(p, "w", encoding="utf-8") as f:
        json.dump(data, f, ensure_ascii=False, indent=2)
        f.write("\n")


def _kuaishou_bg_qr_path(ks_id: str) -> str:
    safe = re.sub(r"[^0-9A-Za-z._-]+", "_", str(ks_id or "").strip()) or "unknown"
    return os.path.join(".state", "kuaishou_bg_qr", f"{safe}.png")


def _run_kuaishou_background_login(
    *,
    profile_dir: str,
    classroom_url: str,
    ks_id: str,
    update_status: callable,
) -> Dict[str, Any]:
    try:
        from playwright.sync_api import sync_playwright
    except Exception as e:
        raise RuntimeError("后台登录快手课堂需要 Playwright") from e

    qr_path = _kuaishou_bg_qr_path(ks_id)
    os.makedirs(os.path.dirname(qr_path), exist_ok=True)

    def _set(progress: int, message: str, **extra: Any) -> None:
        payload = {"progress": int(progress), "message": str(message)}
        payload.update(extra)
        try:
            update_status(payload)
        except Exception:
            pass

    def _is_classroom_url(u: str) -> bool:
        return "kt.kuaishou.com/student-management/offsite-student-management" in (u or "")

    def _find_visible(page: Any, selectors: List[str]) -> Any:
        for sel in selectors:
            try:
                loc = page.locator(sel)
                cnt = loc.count()
            except Exception:
                continue
            for i in range(cnt):
                try:
                    cand = loc.nth(i)
                    if cand.is_visible():
                        return cand
                except Exception:
                    continue
        return None

    with sync_playwright() as p:
        os.makedirs(profile_dir, exist_ok=True)
        ctx = p.chromium.launch_persistent_context(user_data_dir=profile_dir, headless=True)
        try:
            page = ctx.new_page()
            _set(10, "打开快手课堂")
            page.goto(classroom_url, wait_until="domcontentloaded", timeout=60000)
            page.wait_for_timeout(2000)

            cur = page.url or ""
            if _is_classroom_url(cur) and ("passport.kuaishou.com" not in cur) and ("id.kuaishou.com" not in cur):
                try:
                    ctx.storage_state(path=os.path.join(profile_dir, "storage_state.json"))
                except Exception:
                    pass
                _set(100, "已有登录状态", ok=True)
                return {
                    "ok": True,
                    "already_logged_in": True,
                    "finished_at": _now_iso(),
                    "url": cur,
                    "qr_path": qr_path,
                }

            _set(30, "打开扫码登录页")
            switched = False
            for sel in [
                "text=扫码登录",
                "button:has-text('扫码登录')",
                "[role='tab']:has-text('扫码登录')",
                "div:has-text('扫码登录')",
                "span:has-text('扫码登录')",
            ]:
                btn = _find_visible(page, [sel])
                if btn is None:
                    continue
                try:
                    btn.click(timeout=5000)
                    page.wait_for_timeout(1200)
                    switched = True
                    break
                except Exception:
                    continue
            if not switched:
                try:
                    body = (page.locator("body").inner_text(timeout=5000) or "")[:500]
                    _set(35, f"未自动切到扫码页，继续尝试截图。{body}")
                except Exception:
                    pass

            def _save_qr_shot() -> None:
                try:
                    qr_loc = _find_visible(
                        page,
                        [
                            "canvas",
                            "img[alt*='二维码']",
                            "img[src*='qr']",
                            "img[src*='code']",
                            "[class*='qrcode'] img",
                            "[class*='qrcode'] canvas",
                            "[class*='qr'] img",
                            "[class*='qr'] canvas",
                            "[class*='scan'] img",
                            "[class*='scan'] canvas",
                        ],
                    )
                    if qr_loc is not None:
                        qr_loc.screenshot(path=qr_path)
                    else:
                        page.screenshot(path=qr_path, full_page=True)
                except Exception:
                    pass

            _save_qr_shot()
            _set(45, "等待扫码登录", qr_ready=True, qr_path=qr_path)

            deadline = time.time() + 180.0
            while time.time() < deadline:
                cur = page.url or ""
                if _is_classroom_url(cur) and ("passport.kuaishou.com" not in cur) and ("id.kuaishou.com" not in cur):
                    try:
                        page.wait_for_load_state("networkidle", timeout=15000)
                    except Exception:
                        pass
                    page.wait_for_timeout(1000)
                    try:
                        ctx.storage_state(path=os.path.join(profile_dir, "storage_state.json"))
                    except Exception:
                        pass
                    _set(100, "后台登录成功", ok=True, qr_ready=False)
                    return {
                        "ok": True,
                        "already_logged_in": False,
                        "finished_at": _now_iso(),
                        "url": cur,
                        "qr_path": qr_path,
                    }
                _save_qr_shot()
                _set(60, "等待扫码登录", qr_ready=True, qr_path=qr_path)
                page.wait_for_timeout(2000)

            raise RuntimeError("扫码登录超时，请刷新二维码后重试")
        finally:
            try:
                ctx.close()
            except Exception:
                pass


def _save_niu_bg_status(cfg: Dict[str, Any], args: Any, data: Dict[str, Any]) -> None:
    p = _niu_bg_status_json_path(cfg, args)
    os.makedirs(os.path.dirname(p) or ".", exist_ok=True)
    with open(p, "w", encoding="utf-8") as f:
        json.dump(data, f, ensure_ascii=False, indent=2)
        f.write("\n")


def _merge_niu_rows_with_bg_status(rows: List[Dict[str, str]], status_data: Dict[str, Any]) -> List[Dict[str, Any]]:
    out: List[Dict[str, Any]] = []
    items = status_data.get("items") if isinstance(status_data, dict) else {}
    if not isinstance(items, dict):
        items = {}
    for idx, row in enumerate(rows or []):
        one = {"url": str((row or {}).get("url") or "").strip()}
        st = items.get(str(idx))
        if isinstance(st, dict):
            one["bg_login"] = st
        out.append(one)
    return out


def _run_niu_background_login(
    *,
    url: str,
    profile_dir: str,
    phone: str,
    password: str,
    update_status: callable,
) -> Dict[str, Any]:
    try:
        from playwright.sync_api import sync_playwright
    except Exception as e:
        raise RuntimeError("后台登录金牛需要 Playwright") from e

    def _set(progress: int, message: str) -> None:
        try:
            update_status(progress, message)
        except Exception:
            pass

    _set(5, "启动后台浏览器")
    with sync_playwright() as p:
        os.makedirs(profile_dir, exist_ok=True)
        ctx = p.chromium.launch_persistent_context(user_data_dir=profile_dir, headless=True)
        try:
            page = ctx.new_page()
            _set(15, "打开金牛页面")
            page.goto(url, wait_until="domcontentloaded", timeout=60000)
            page.wait_for_timeout(1500)

            def _body_text() -> str:
                try:
                    return (page.locator("body").inner_text(timeout=5000) or "").strip()
                except Exception:
                    return ""

            def _visible(selectors: List[str]):
                for sel in selectors:
                    try:
                        loc = page.locator(sel)
                        cnt = loc.count()
                    except Exception:
                        continue
                    for i in range(cnt):
                        try:
                            cand = loc.nth(i)
                            if cand.is_visible():
                                return cand
                        except Exception:
                            continue
                return None

            phone_input = _visible([
                "input[placeholder='邮箱/快手号绑定的手机号']",
                "input[placeholder='请输入手机号']",
            ])
            if phone_input is None:
                _set(25, "展开登录面板")
                for sel in ["button:has-text('立即登录')", "text=立即登录", "button:has-text('登录')", "text=登录"]:
                    btn = _visible([sel])
                    if btn is None:
                        continue
                    try:
                        btn.click(timeout=5000)
                        page.wait_for_timeout(1000)
                        phone_input = _visible([
                            "input[placeholder='邮箱/快手号绑定的手机号']",
                            "input[placeholder='请输入手机号']",
                        ])
                        if phone_input is not None:
                            break
                    except Exception:
                        continue

            if phone_input is None:
                raise RuntimeError("未找到金牛登录账号输入框")

            password_input = _visible(["input[placeholder='请输入密码']", "input[type='password']"])
            if password_input is None:
                raise RuntimeError("未找到金牛登录密码输入框")

            _set(40, "填写账号密码")
            phone_input.fill(phone)
            password_input.fill(password)

            try:
                agree = _visible(["input[type='checkbox']"])
                if agree is not None and (not agree.is_checked()):
                    agree.check(force=True)
            except Exception:
                pass

            _set(60, "提交登录请求")
            clicked = False
            for sel in ["button:has-text('登录')", "text=登录"]:
                btn = _visible([sel])
                if btn is None:
                    continue
                try:
                    btn.click(timeout=5000)
                    clicked = True
                    break
                except Exception:
                    continue
            if not clicked:
                raise RuntimeError("未找到可点击的登录按钮")

            _set(75, "等待后台登录完成")
            deadline = time.time() + 45.0
            while time.time() < deadline:
                cur = page.url or ""
                if ("/reportV2/commonReport" in cur) and ("/welcome" not in cur):
                    break
                page.wait_for_timeout(1000)
            try:
                page.wait_for_load_state("networkidle", timeout=15000)
            except Exception:
                pass
            page.wait_for_timeout(1500)

            final_url = page.url or ""
            body = _body_text()
            if ("/reportV2/commonReport" not in final_url) or ("/welcome" in final_url):
                raise RuntimeError(f"后台登录后仍未进入列表页。url={final_url} body={body[:160]}")

            _set(90, "验证登录状态")
            has_report_marker = False
            try:
                markers = ["直播ID", "直播中", "最近7天", "条/页", "每页"]
                joined = body[:4000]
                has_report_marker = any(m in joined for m in markers)
            except Exception:
                has_report_marker = False

            try:
                ctx.storage_state(path=os.path.join(profile_dir, "storage_state.json"))
            except Exception:
                pass

            _set(100, "后台登录完成")
            return {
                "ok": True,
                "url": final_url,
                "profile_dir": profile_dir,
                "verified": bool(has_report_marker),
                "finished_at": _now_iso(),
            }
        finally:
            try:
                ctx.close()
            except Exception:
                pass


def _copy_table_cache_path() -> str:
    return os.path.join(".state", "copy_table_cache.json")


def _load_copy_table_cache() -> Dict[str, Any]:
    p = _copy_table_cache_path()
    try:
        if os.path.exists(p):
            with open(p, "r", encoding="utf-8") as f:
                data = json.load(f)
            return data if isinstance(data, dict) else {}
        return {}
    except Exception:
        return {}


def _save_copy_table_cache(data: Dict[str, Any]) -> None:
    p = _copy_table_cache_path()
    os.makedirs(os.path.dirname(p) or ".", exist_ok=True)
    with open(p, "w", encoding="utf-8") as f:
        json.dump(data, f, ensure_ascii=False, indent=2)
        f.write("\n")


def _set_copy_table_item_copied(key: str, copied: bool) -> bool:
    key = (key or "").strip()
    if not key:
        return False
    data = _load_copy_table_cache()
    items = data.get("items")
    if not isinstance(items, list):
        return False
    changed = False
    for it in items:
        if not isinstance(it, dict):
            continue
        if str(it.get("key") or "") == key:
            it["copied"] = bool(copied)
            changed = True
            break
    if changed:
        data["items"] = items
        data["updated_at"] = _now_ts()
        _save_copy_table_cache(data)
    return changed


def _clear_copy_table_cache() -> None:
    try:
        p = _copy_table_cache_path()
        if os.path.exists(p):
            os.remove(p)
    except Exception:
        pass


def _load_user_contact_dept(cfg: Dict[str, Any], args: Any) -> str:
    p = _user_contact_cfg_path(cfg, args)
    try:
        with open(p, "r", encoding="utf-8") as f:
            data = json.load(f)
        v = (data.get("dept") or "").strip()
        return v
    except FileNotFoundError:
        return ""
    except Exception:
        return ""


def _save_user_contact_dept(cfg: Dict[str, Any], args: Any, dept: str) -> None:
    p = _user_contact_cfg_path(cfg, args)
    os.makedirs(os.path.dirname(p) or ".", exist_ok=True)
    with open(p, "w", encoding="utf-8") as f:
        json.dump({"dept": (dept or "").strip()}, f, ensure_ascii=False, indent=2)


def _load_wenzong_user_contact_dept(cfg: Dict[str, Any], args: Any) -> str:
    p = _wenzong_user_contact_cfg_path(cfg, args)
    try:
        with open(p, "r", encoding="utf-8") as f:
            data = json.load(f)
        v = (data.get("dept") or "").strip()
        return v
    except FileNotFoundError:
        return ""
    except Exception:
        return ""


def _save_wenzong_user_contact_dept(cfg: Dict[str, Any], args: Any, dept: str) -> None:
    p = _wenzong_user_contact_cfg_path(cfg, args)
    os.makedirs(os.path.dirname(p) or ".", exist_ok=True)
    with open(p, "w", encoding="utf-8") as f:
        json.dump({"dept": (dept or "").strip()}, f, ensure_ascii=False, indent=2)


def _next_user_contact_seq(cfg: Dict[str, Any], args: Any) -> int:
    p = _user_contact_seq_path(cfg, args)
    try:
        if os.path.exists(p):
            with open(p, "r", encoding="utf-8") as f:
                data = json.load(f)
            v = int(data.get("seq", 0))
        else:
            v = 0
    except Exception:
        v = 0
    v += 1
    try:
        os.makedirs(os.path.dirname(p) or ".", exist_ok=True)
        with open(p, "w", encoding="utf-8") as f:
            json.dump({"seq": v}, f, ensure_ascii=False, indent=2)
    except Exception:
        pass
    return v


def _next_wenzong_user_contact_seq(cfg: Dict[str, Any], args: Any) -> int:
    p = _wenzong_user_contact_seq_path(cfg, args)
    try:
        if os.path.exists(p):
            with open(p, "r", encoding="utf-8") as f:
                data = json.load(f)
            v = int(data.get("seq", 0))
        else:
            v = 0
    except Exception:
        v = 0
    v += 1
    try:
        os.makedirs(os.path.dirname(p) or ".", exist_ok=True)
        with open(p, "w", encoding="utf-8") as f:
            json.dump({"seq": v}, f, ensure_ascii=False, indent=2)
    except Exception:
        pass
    return v


def _export_user_contact_xlsx_to_dir(
    cfg: Dict[str, Any],
    args: Any,
    rows: List[List[str]],
    *,
    export_dir: str,
) -> Dict[str, Any]:
    if not export_dir:
        return {"ok": False, "error": "export_dir_not_configured"}
    if export_dir.lower().startswith("smb://"):
        return {
            "ok": False,
            "error": "export_dir_is_smb_url",
            "export_dir": export_dir,
        }

    dept = _load_user_contact_dept(cfg, args)
    rows2 = _apply_dept(rows, dept)
    added_count = max(0, len(rows2 or []))
    name = f"新增{added_count}、{_now_ts()}.xlsx"
    try:
        payload = _xlsx_bytes_from_rows_template(cfg, args, rows2)
    except Exception as e:
        print(f"[web] template export failed, fallback to simple workbook: {e}", file=sys.stderr)
        payload = _xlsx_bytes_from_rows(rows2)

    try:
        os.makedirs(export_dir, exist_ok=True)
        out_path = os.path.join(export_dir, name)
        with open(out_path, "wb") as f:
            f.write(payload)
        print(f"[web] user_contact xlsx exported to: {out_path}", file=sys.stderr)
        return {"ok": True, "name": name, "out_path": out_path}
    except Exception as e:
        print(f"[web] failed to write user_contact xlsx to export_dir: {e}", file=sys.stderr)
        return {"ok": False, "error": str(e), "name": name}


class AccountRunner:
    def __init__(
        self,
        *,
        repo_dir: str,
        config_path: str,
        account: str,
        interval_seconds: int,
        headless: bool,
        global_run_lock: threading.Lock,
        log_dir: str,
        exports_root: str,
        sync_logs_dir: str,
        keep_exports: int,
        keep_sync_logs: int,
        keep_web_logs: int,
    ):
        self.repo_dir = repo_dir
        self.config_path = config_path
        self.account = account
        self.interval_seconds = interval_seconds
        self.headless = headless
        self.global_run_lock = global_run_lock
        self.log_dir = log_dir
        self.exports_root = exports_root
        self.sync_logs_dir = sync_logs_dir
        self.keep_exports = keep_exports
        self.keep_sync_logs = keep_sync_logs
        self.keep_web_logs = keep_web_logs

        self._live_id_lock = threading.Lock()
        self.live_id: str = ""

        self._thread: Optional[threading.Thread] = None
        self._stop = threading.Event()
        self._run_lock = threading.Lock()

        self.last_started_at: Optional[float] = None
        self.last_finished_at: Optional[float] = None
        self.last_returncode: Optional[int] = None
        self.last_error: str = ""
        self.last_log_path: str = ""

        self._next_run_lock = threading.Lock()
        self.next_run_at: Optional[float] = None

    def is_running(self) -> bool:
        return self._thread is not None and self._thread.is_alive() and not self._stop.is_set()

    def start(self) -> None:
        if self.is_running():
            return
        self._stop.clear()
        t = threading.Thread(target=self._loop, name=f"runner:{self.account}", daemon=True)
        self._thread = t
        t.start()

    def stop(self) -> None:
        self._stop.set()

    def set_live_id(self, live_id: str) -> None:
        with self._live_id_lock:
            self.live_id = (live_id or "").strip()

    def get_live_id(self) -> str:
        with self._live_id_lock:
            return self.live_id

    def _loop(self) -> None:
        while not self._stop.is_set():
            self.run_once()
            next_run_at = time.time() + float(self.interval_seconds)
            with self._next_run_lock:
                self.next_run_at = next_run_at
            while not self._stop.is_set():
                remain = int(round(next_run_at - time.time()))
                if remain <= 0:
                    break
                if self._stop.wait(1.0):
                    break
            if self._stop.is_set():
                break
        with self._next_run_lock:
            self.next_run_at = None

    def run_once(self) -> None:
        if not self._run_lock.acquire(blocking=False):
            return
        acquired_global = False
        try:
            self.last_started_at = time.time()
            self.last_error = ""

            if not self.global_run_lock.acquire(blocking=False):
                self.last_returncode = -2
                self.last_finished_at = time.time()
                self.last_error = "busy: another account is running"
                return
            acquired_global = True

            cmd = [
                sys.executable,
                "run_pipeline.py",
                "--config",
                self.config_path,
                "--account",
                self.account,
            ]
            live_id = self.get_live_id()
            if live_id:
                cmd.extend(["--live-id", live_id])
            if self.headless:
                cmd.append("--headless")

            proc = subprocess.run(cmd, cwd=self.repo_dir, capture_output=True, text=True)
            self.last_returncode = proc.returncode
            self.last_finished_at = time.time()

            os.makedirs(self.log_dir, exist_ok=True)
            self.last_log_path = os.path.join(self.log_dir, f"{self.account}.log")
            with open(self.last_log_path, "a", encoding="utf-8") as f:
                f.write("\n" + "=" * 72 + "\n")
                f.write(f"started_at={_fmt_ts(self.last_started_at)}\n")
                f.write(f"finished_at={_fmt_ts(self.last_finished_at)}\n")
                f.write(f"returncode={proc.returncode}\n")
                f.write(f"cmd={' '.join(cmd)}\n")
                if proc.stdout:
                    f.write("\n[stdout]\n")
                    f.write(proc.stdout)
                    if not proc.stdout.endswith("\n"):
                        f.write("\n")
                if proc.stderr:
                    f.write("\n[stderr]\n")
                    f.write(proc.stderr)
                    if not proc.stderr.endswith("\n"):
                        f.write("\n")

            if proc.returncode != 0:
                err = (proc.stderr or proc.stdout or "").strip()
                self.last_error = err[:2000]

            export_dir = os.path.join(self.exports_root, self.account)
            _cleanup_keep_latest_files(export_dir, keep=self.keep_exports)
            _cleanup_keep_latest_files(self.sync_logs_dir, keep=self.keep_sync_logs, exts=[".log"])
            _cleanup_keep_latest_files(self.log_dir, keep=self.keep_web_logs, exts=[".log"])
        except Exception as e:
            self.last_returncode = -1
            self.last_finished_at = time.time()
            self.last_error = str(e)
        finally:
            if acquired_global:
                try:
                    self.global_run_lock.release()
                except Exception:
                    pass
            self._run_lock.release()


class GlobalRunner:
    def __init__(
        self,
        *,
        repo_dir: str,
        config_path: str,
        interval_seconds: int,
        headless: bool,
        global_run_lock: threading.Lock,
        log_dir: str,
        sync_logs_dir: str,
        keep_sync_logs: int,
        keep_web_logs: int,
    ):
        self.repo_dir = repo_dir
        self.config_path = config_path
        self.interval_seconds = interval_seconds
        self.headless = headless
        self.global_run_lock = global_run_lock
        self.log_dir = log_dir
        self.sync_logs_dir = sync_logs_dir
        self.keep_sync_logs = keep_sync_logs
        self.keep_web_logs = keep_web_logs

        self._thread: Optional[threading.Thread] = None
        self._stop = threading.Event()
        self._run_lock = threading.Lock()

        self.last_started_at: Optional[float] = None
        self.last_finished_at: Optional[float] = None
        self.last_returncode: Optional[int] = None
        self.last_error: str = ""
        self.last_log_path: str = ""

        self._next_run_lock = threading.Lock()
        self.next_run_at: Optional[float] = None

        self.night_sleep_enabled: bool = False
        self._sleeping: bool = False

    def is_running(self) -> bool:
        return self._thread is not None and self._thread.is_alive() and not self._stop.is_set()

    def is_in_sleep_window(self) -> bool:
        """检查当前是否在夜间休眠时间窗口 (04:00 ~ 10:00)"""
        if not self.night_sleep_enabled:
            return False
        now = datetime.now()
        return 4 <= now.hour < 10

    def is_sleeping(self) -> bool:
        return self._sleeping

    def start(self) -> None:
        if self.is_running():
            return
        self._stop.clear()
        t = threading.Thread(target=self._loop, name="runner:global", daemon=True)
        self._thread = t
        t.start()

    def stop(self) -> None:
        self._stop.set()

    def set_interval(self, interval_seconds: int) -> None:
        """动态设置间隔时间"""
        self.interval_seconds = max(1, int(interval_seconds))

    def get_interval_seconds(self) -> int:
        return int(self.interval_seconds)

    def get_next_run_in_seconds(self) -> Optional[int]:
        with self._next_run_lock:
            nra = self.next_run_at
        if not nra:
            return None
        try:
            return max(0, int(round(nra - time.time())))
        except Exception:
            return None

    def _loop(self) -> None:
        while not self._stop.is_set():
            if self.is_in_sleep_window():
                self._sleeping = True
                self.last_error = "夜间休眠中 (04:00~10:00)"
                with self._next_run_lock:
                    self.next_run_at = None
                if self._stop.wait(30.0):
                    break
                continue
            self._sleeping = False
            self.run_once()
            next_run_at = time.time() + float(self.interval_seconds)
            with self._next_run_lock:
                self.next_run_at = next_run_at
            while not self._stop.is_set():
                remain = int(round(next_run_at - time.time()))
                if remain <= 0:
                    break
                if self._stop.wait(1.0):
                    break
            if self._stop.is_set():
                break
        self._sleeping = False
        with self._next_run_lock:
            self.next_run_at = None

    def run_once(self) -> None:
        if not self._run_lock.acquire(blocking=False):
            return
        acquired_global = False
        try:
            self.last_started_at = time.time()
            self.last_error = ""

            if not self.global_run_lock.acquire(blocking=False):
                self.last_returncode = -2
                self.last_finished_at = time.time()
                self.last_error = "busy: another run is running"
                return
            acquired_global = True

            cmd = [
                sys.executable,
                "run_pipeline.py",
                "--config",
                self.config_path,
                "--all-accounts",
            ]
            if self.headless:
                cmd.append("--headless")

            proc = subprocess.run(cmd, cwd=self.repo_dir, capture_output=True, text=True)
            self.last_returncode = proc.returncode
            self.last_finished_at = time.time()

            os.makedirs(self.log_dir, exist_ok=True)
            self.last_log_path = os.path.join(self.log_dir, "global.log")
            _rotate_log_file_if_needed(
                self.last_log_path,
                max_bytes=20 * 1024 * 1024,
                keep=self.keep_web_logs,
            )
            with open(self.last_log_path, "a", encoding="utf-8") as f:
                f.write("\n" + "=" * 72 + "\n")
                f.write(f"started_at={_fmt_ts(self.last_started_at)}\n")
                f.write(f"finished_at={_fmt_ts(self.last_finished_at)}\n")
                f.write(f"returncode={proc.returncode}\n")
                f.write(f"cmd={' '.join(cmd)}\n")
                if proc.stdout:
                    f.write("\n[stdout]\n")
                    f.write(proc.stdout)
                    if not proc.stdout.endswith("\n"):
                        f.write("\n")
                if proc.stderr:
                    f.write("\n[stderr]\n")
                    f.write(proc.stderr)
                    if not proc.stderr.endswith("\n"):
                        f.write("\n")

            if proc.returncode != 0:
                err = (proc.stderr or proc.stdout or "").strip()
                self.last_error = err[:2000]

            # Keep browser profiles from growing without bound between manual cleanups.
            try:
                _clean_browser_profile_caches(
                    [
                        '.state/kuaishou_profiles',
                        '.state/niu_chrome_profile',
                        '.state/feishu_profile',
                        '.state/niu_profile',
                        '.state/kuaishou_chromium',
                        '.state/niu_pw_profiles',
                    ]
                )
            except Exception:
                pass

            _cleanup_keep_latest_files(self.sync_logs_dir, keep=self.keep_sync_logs, exts=[".log"])
            _cleanup_keep_latest_files(self.log_dir, keep=self.keep_web_logs, exts=[".log"])
        except Exception as e:
            self.last_returncode = -1
            self.last_finished_at = time.time()
            self.last_error = str(e)
        finally:
            if acquired_global:
                try:
                    self.global_run_lock.release()
                except Exception:
                    pass
            self._run_lock.release()


def _now_epoch() -> float:
    return time.time()


def _fmt_ts(ts: Optional[float]) -> str:
    if not ts:
        return ""
    try:
        return time.strftime("%Y-%m-%d %H:%M:%S", time.localtime(ts))
    except Exception:
        return str(ts)


def _rotate_log_file_if_needed(path: str, *, max_bytes: int, keep: int) -> None:
    try:
        if max_bytes <= 0:
            return
        if not path or (not os.path.exists(path)):
            return
        if os.path.getsize(path) < int(max_bytes):
            return
        ts = time.strftime("%Y%m%d_%H%M%S", time.localtime())
        base_dir = os.path.dirname(os.path.abspath(path))
        base_name = os.path.basename(path)
        name_root, name_ext = os.path.splitext(base_name)
        rotated = os.path.join(base_dir, f"{name_root}.{ts}{name_ext or '.log'}")
        try:
            os.replace(path, rotated)
        except FileNotFoundError:
            return
        except Exception:
            return
        try:
            _cleanup_keep_latest_files(base_dir, keep=int(keep) if keep else 20, exts=[".log"])
        except Exception:
            pass
    except Exception:
        return


def build_handler(
    *,
    token: str,
    accounts: List[str],
    runners: Dict[str, AccountRunner],
    global_runner: GlobalRunner,
    cfg: Dict[str, Any],
    args: Any,
    reload_callback: callable,
) -> type:
    niu_bg_state_lock = threading.Lock()
    niu_bg_runtime: Dict[str, Dict[str, Any]] = {}
    kuaishou_bg_state_lock = threading.Lock()
    kuaishou_bg_runtime: Dict[str, Dict[str, Any]] = {}

    def _get_niu_status(idx: int) -> Dict[str, Any]:
        key = str(idx)
        with niu_bg_state_lock:
            rt = niu_bg_runtime.get(key)
            if isinstance(rt, dict):
                return dict(rt)
        saved = _load_niu_bg_status(cfg, args)
        items = saved.get("items") if isinstance(saved, dict) else {}
        if isinstance(items, dict) and isinstance(items.get(key), dict):
            return dict(items.get(key) or {})
        return {"ok": False, "running": False, "progress": 0, "message": "未登录"}

    def _set_niu_status(idx: int, patch: Dict[str, Any], *, persist: bool = False) -> Dict[str, Any]:
        key = str(idx)
        with niu_bg_state_lock:
            cur = dict(niu_bg_runtime.get(key) or {})
            cur.update(patch or {})
            niu_bg_runtime[key] = cur
        if persist:
            saved = _load_niu_bg_status(cfg, args)
            items = saved.get("items")
            if not isinstance(items, dict):
                items = {}
            items[key] = {k: v for k, v in cur.items() if k != "running"}
            saved["items"] = items
            saved["updated_at"] = _now_iso()
            _save_niu_bg_status(cfg, args, saved)
        return cur

    def _kuaishou_bg_key(scope: str, ks_id: str) -> str:
        scope_s = str(scope or "").strip().lower() or "default"
        ks_s = str(ks_id or "").strip()
        return f"{scope_s}:{ks_s}"

    def _get_kuaishou_status(scope: str, ks_id: str) -> Dict[str, Any]:
        key = _kuaishou_bg_key(scope, ks_id)
        with kuaishou_bg_state_lock:
            rt = kuaishou_bg_runtime.get(key)
            if isinstance(rt, dict):
                return dict(rt)
        saved = _load_kuaishou_bg_status(cfg, args)
        items = saved.get("items") if isinstance(saved, dict) else {}
        if isinstance(items, dict) and isinstance(items.get(key), dict):
            return dict(items.get(key) or {})
        return {"ok": False, "running": False, "progress": 0, "message": "未登录"}

    def _set_kuaishou_status(scope: str, ks_id: str, patch: Dict[str, Any], *, persist: bool = False) -> Dict[str, Any]:
        key = _kuaishou_bg_key(scope, ks_id)
        with kuaishou_bg_state_lock:
            cur = dict(kuaishou_bg_runtime.get(key) or {})
            cur.update(patch or {})
            kuaishou_bg_runtime[key] = cur
        if persist:
            saved = _load_kuaishou_bg_status(cfg, args)
            items = saved.get("items")
            if not isinstance(items, dict):
                items = {}
            items[key] = {k: v for k, v in cur.items() if k != "running"}
            saved["items"] = items
            saved["updated_at"] = _now_iso()
            _save_kuaishou_bg_status(cfg, args, saved)
        return cur


    index_html = """<!doctype html>
<html lang=\"zh\">
<head>
  <meta charset=\"utf-8\" />
  <meta name=\"viewport\" content=\"width=device-width, initial-scale=1\" />
  <title>OpenClaw 控制台</title>
  <style>
    body { font-family: -apple-system, BlinkMacSystemFont, Segoe UI, Roboto, Helvetica, Arial; margin: 24px; background: #0b0f17; color: #e7eefc; }
    .card { background: #111827; border: 1px solid #1f2a44; border-radius: 12px; padding: 16px; margin-bottom: 12px; }
    .row { display: flex; align-items: center; justify-content: space-between; gap: 12px; }
    .muted { color: #9fb0d0; font-size: 12px; }
    .title { font-size: 18px; font-weight: 700; }
    button { border: 1px solid #2a3a5f; background: #16213a; color: #e7eefc; padding: 8px 12px; border-radius: 10px; cursor: pointer; }
    button:hover { background: #1b2b4d; }
    button.danger { border-color: #6b1b1b; background: #2a1010; }
    button.danger:hover { background: #3a1414; }
    .pill { padding: 4px 8px; border-radius: 999px; font-size: 12px; border: 1px solid #2a3a5f; }
    .on { background: #0f2a1c; border-color: #1d6b3d; }
    .off { background: #221427; border-color: #5a2b6a; }
    pre { white-space: pre-wrap; word-break: break-word; background: #0b1220; border: 1px solid #1f2a44; border-radius: 12px; padding: 12px; }
    .tabs { display: flex; gap: 8px; margin-bottom: 16px; flex-wrap: wrap; }
    .tab { padding: 10px 20px; background: #111827; border: 1px solid #1f2a44; border-radius: 8px; cursor: pointer; }
    .tab.active { background: #16213a; border-color: #2a3a5f; }
    .tab-content { display: none; }
    .tab-content.active { display: block; }
    table { width: 100%; border-collapse: collapse; }
    th, td { padding: 10px; text-align: left; border-bottom: 1px solid #1f2a44; }
    th { background: #0b1220; font-weight: 600; }
    input[type="text"] { background: #0b1220; color: #e7eefc; border: 1px solid #2a3a5f; padding: 6px 10px; border-radius: 6px; width: 100%; }
    input[type="number"] { background: #0b1220; color: #e7eefc; border: 1px solid #2a3a5f; padding: 6px 10px; border-radius: 6px; width: 80px; }
    .btn-small { padding: 4px 8px; font-size: 12px; }
    .inline-group { display: flex; align-items: center; gap: 8px; flex-wrap: wrap; }
    .modal-mask { position: fixed; inset: 0; background: rgba(2, 6, 23, 0.72); display: none; align-items: center; justify-content: center; padding: 20px; z-index: 9999; }
    .modal-mask.show { display: flex; }
    .modal { width: min(480px, 100%); background: #111827; border: 1px solid #2a3a5f; border-radius: 14px; padding: 18px; }
    .modal-wide { width: min(920px, 96vw); }
    .modal h3 { margin: 0 0 12px; font-size: 18px; }
    .field { margin-bottom: 12px; }
    .field label { display: block; margin-bottom: 6px; color: #9fb0d0; font-size: 12px; }
    .field input[type="password"] { background: #0b1220; color: #e7eefc; border: 1px solid #2a3a5f; padding: 6px 10px; border-radius: 6px; width: 100%; }
    .progress { width: 100%; height: 10px; background: #0b1220; border: 1px solid #1f2a44; border-radius: 999px; overflow: hidden; margin-top: 10px; }
    .progress-bar { height: 100%; width: 0%; background: linear-gradient(90deg, #2dd4bf, #60a5fa); transition: width 0.25s ease; }
    .status-ok { color: #86efac; }
    .status-warn { color: #fcd34d; }
    .status-err { color: #fca5a5; }
  </style>
</head>
<body>
  <div class="card">
    <div class="row">
      <div>
        <div class="title">OpenClaw 控制台</div>
        <div class="muted">server_version: %SERVER_VERSION%</div>
      </div>

    </div>
  </div>

  <div class="tabs">
    <div class="tab active" onclick="switchTab('global')">运行控制</div>
    <div class="tab" onclick="switchTab('mapping')">主播映射表</div>
    <div class="tab" onclick="switchTab('mapping_wenzong')">WenZong 映射表</div>
    <div class="tab" onclick="switchTab('niu')">金牛表</div>
    <div class="tab" onclick="switchTab('copy')">复制数据表</div>
  </div>

  <div id="tab-global" class="tab-content active">
    <div class="card">
      <div class="row">
        <div class="muted">总开关：开启后自动执行（run_pipeline.py --all-accounts）</div>
        <div>
          <button type="button" onclick="refresh()">刷新</button>
          <button type="button" onclick="startGlobal()">Start</button>
          <button type="button" class="danger" onclick="stopGlobal()">Stop</button>
          <button type="button" onclick="runOnceGlobal()">Run once</button>
        </div>
      </div>
      <div style="margin-top:10px" class="muted" id="global-status"></div>
      <div style="margin-top:6px" class="muted" id="next-run"></div>
      <div style="margin-top:10px" id="global-error"></div>
      <div style="margin-top:14px" class="muted">金牛实时状态</div>
      <div style="margin-top:8px" id="niu-runtime-status" class="muted"></div>
      <div style="margin-top:14px" class="muted">快手课堂实时状态</div>
      <div style="margin-top:8px" id="kuaishou-runtime-status" class="muted"></div>
      <div style="margin-top:14px" class="muted">云文档与导出实时状态</div>
      <div style="margin-top:8px" id="feishu-runtime-status" class="muted"></div>
    </div>

    <div class="card">
      <div class="row">
        <div class="muted">WenZong 流程：主流程完成后再执行一轮（独立云文档/独立主播映射表；不写投放信息表）</div>
        <div class="inline-group">
          <label class="muted"><input type="checkbox" id="wenzong-enabled" /> 启用</label>
          <span class="muted">WenZong部门</span>
          <input type="text" id="wz-uc-dept" placeholder="可编辑并保存" style="width: 240px;" />
          <button type="button" onclick="saveWenzongUserContactDept()">保存</button>
        </div>
      </div>
    </div>

    <div class="card">
      <div class="row">
        <div class="muted">用户对接信息表：维护部门（可保存），导入/保存数据后自动生成xlsx到导出目录</div>
        <div class="inline-group">
          <span class="muted">部门</span>
          <input type="text" id="uc-dept" placeholder="可编辑并保存" style="width: 240px;" />
          <button type="button" onclick="saveUserContactDept()">保存部门</button>
          <input type="file" id="uc-file" accept=".xlsx" style="display:none" onchange="importUserContactXlsx(event)" />
          <button type="button" onclick="document.getElementById('uc-file').click()">导入xlsx</button>
        </div>
      </div>
    </div>

    <div class="card">
      <div class="row">
        <div class="muted">执行间隔设置</div>
        <div class="inline-group">
          <span class="muted">每</span>
          <input type="number" id="interval-input" min="1" max="3600" value="120" />
          <span class="muted">秒执行一次</span>
          <button type="button" onclick="saveInterval()">保存间隔</button>
        </div>
      </div>
    </div>

    <div class="card">
      <div class="row">
        <div>
          <div class="muted">夜间休眠：开启后凌晨 4:00 ~ 10:00 暂停脚本循环</div>
          <div class="muted" id="night-sleep-status" style="margin-top:4px"></div>
        </div>
        <div class="inline-group">
          <label style="position:relative;display:inline-block;width:48px;height:26px;margin:0">
            <input type="checkbox" id="night-sleep-toggle" onchange="toggleNightSleep()" style="opacity:0;width:0;height:0" />
            <span style="position:absolute;cursor:pointer;inset:0;background:#2a1010;border:1px solid #6b1b1b;border-radius:26px;transition:.3s"></span>
            <span id="night-sleep-knob" style="position:absolute;left:3px;top:3px;width:20px;height:20px;background:#e7eefc;border-radius:50%;transition:.3s"></span>
          </label>
        </div>
      </div>
    </div>

    <div class="card">
      <div class="row">
        <div class="muted">浏览器缓存清理（保留登录状态）</div>
        <div>
          <button type="button" onclick="cleanCache()">清理缓存</button>
          <span class="muted" id="cache-status"></span>
        </div>
      </div>
    </div>
  </div>

  <div id="tab-mapping" class="tab-content">
    <div class="card">
      <div class="row">
        <div class="muted">当前使用的主播账号</div>
        <div>
          <button type="button" onclick="addMappingRow('active')">新增行</button>
          <button type="button" onclick="saveMappingTables()">保存</button>
          <button type="button" onclick="loadMappingTables()">刷新</button>
        </div>
      </div>
    </div>

    <div class="card">
      <table id="mapping-table-active">
        <thead>
          <tr>
            <th>直播账号</th>
            <th>快手ID</th>
            <th>手机号码</th>
            <th>密码</th>
            <th>主播</th>
            <th>快手课堂</th>
            <th>操作</th>
          </tr>
        </thead>
        <tbody id="mapping-tbody-active">
        </tbody>
      </table>
    </div>

    <div class="card" style="margin-top: 24px;">
      <div class="row">
        <div class="muted">暂时不用的主播账号</div>
        <div>
          <button type="button" onclick="addMappingRow('inactive')">新增行</button>
        </div>
      </div>
    </div>

    <div class="card">
      <table id="mapping-table-inactive">
        <thead>
          <tr>
            <th>直播账号</th>
            <th>快手ID</th>
            <th>手机号码</th>
            <th>密码</th>
            <th>主播</th>
            <th>快手课堂</th>
            <th>操作</th>
          </tr>
        </thead>
        <tbody id="mapping-tbody-inactive">
        </tbody>
      </table>
    </div>
  </div>

  <div id="tab-mapping_wenzong" class="tab-content">
    <div class="card">
      <div class="row">
        <div class="muted">WenZong 使用的主播账号（独立文件）</div>
        <div>
          <button type="button" onclick="addMappingRowWenzong('active')">新增行</button>
          <button type="button" onclick="saveMappingTablesWenzong()">保存</button>
          <button type="button" onclick="loadMappingTablesWenzong()">刷新</button>
        </div>
      </div>
    </div>

    <div class="card">
      <table id="mapping-table-active-wenzong">
        <thead>
          <tr>
            <th>直播账号</th>
            <th>快手ID</th>
            <th>手机号码</th>
            <th>密码</th>
            <th>主播</th>
            <th>快手课堂</th>
            <th>操作</th>
          </tr>
        </thead>
        <tbody id="mapping-tbody-active-wenzong">
        </tbody>
      </table>
    </div>

    <div class="card" style="margin-top: 24px;">
      <div class="row">
        <div class="muted">WenZong 暂时不用的主播账号</div>
        <div>
          <button type="button" onclick="addMappingRowWenzong('inactive')">新增行</button>
        </div>
      </div>
    </div>

    <div class="card">
      <table id="mapping-table-inactive-wenzong">
        <thead>
          <tr>
            <th>直播账号</th>
            <th>快手ID</th>
            <th>手机号码</th>
            <th>密码</th>
            <th>主播</th>
            <th>快手课堂</th>
            <th>操作</th>
          </tr>
        </thead>
        <tbody id="mapping-tbody-inactive-wenzong">
        </tbody>
      </table>
    </div>
  </div>

  <div id="tab-niu" class="tab-content">
    <div class="card">
      <div class="row">
        <div class="muted">金牛表：维护多个金牛网址，点击“打开”分别登录不同账号（每行使用独立浏览器 profile）</div>
        <div>
          <button type="button" onclick="addNiuRow()">新增行</button>
          <button type="button" onclick="saveNiuTable()">保存</button>
          <button type="button" onclick="loadNiuTable()">刷新</button>
        </div>
      </div>
    </div>

    <div class="card">
      <table id="niu-table">
        <thead>
          <tr>
            <th>金牛网址</th>
            <th>后台状态</th>
            <th>操作</th>
          </tr>
        </thead>
        <tbody id="niu-tbody"></tbody>
      </table>
    </div>
  </div>

  <div id="niu-login-modal" class="modal-mask">
    <div class="modal">
      <h3>金牛后台登录</h3>
      <div class="muted" id="niu-login-modal-url"></div>
      <div class="field">
        <label for="niu-login-phone">账号</label>
        <input id="niu-login-phone" type="text" placeholder="请输入金牛账号" />
      </div>
      <div class="field">
        <label for="niu-login-password">密码</label>
        <input id="niu-login-password" type="password" placeholder="请输入金牛密码" />
      </div>
      <div class="inline-group">
        <button type="button" id="niu-login-submit" onclick="submitNiuBackgroundLogin()">开始后台登录</button>
        <button type="button" class="danger" onclick="closeNiuLoginModal()">关闭</button>
      </div>
      <div class="progress">
        <div id="niu-login-progress-bar" class="progress-bar"></div>
      </div>
      <div id="niu-login-progress-text" class="muted" style="margin-top:8px"></div>
    </div>
  </div>

  <div id="kuaishou-login-modal" class="modal-mask">
    <div class="modal modal-wide">
      <h3>快手课堂后台登录</h3>
      <div class="muted" id="kuaishou-login-modal-title"></div>
      <div class="inline-group" style="margin-top:10px">
        <button type="button" id="kuaishou-login-submit" onclick="submitKuaishouBackgroundLogin()">后台开始登录</button>
        <button type="button" class="danger" onclick="closeKuaishouLoginModal()">关闭</button>
      </div>
      <div class="progress">
        <div id="kuaishou-login-progress-bar" class="progress-bar"></div>
      </div>
      <div id="kuaishou-login-progress-text" class="muted" style="margin-top:8px"></div>
      <div style="margin-top:14px">
        <img id="kuaishou-login-qr" alt="快手扫码二维码" style="max-width:100%; max-height:70vh; border-radius:12px; border:1px solid #1f2a44; display:none;" />
        <div id="kuaishou-login-qr-empty" class="muted" style="margin-top:8px">如需扫码登录，二维码会显示在这里</div>
      </div>
    </div>
  </div>

  <div id="tab-copy" class="tab-content">
    <div class="card">
      <div class="row">
        <div class="muted">复制数据表：来源于同步到飞书“用户对接信息表”的新增行，点击复制后会标记为已复制（本地缓存）</div>
        <div>
          <button type="button" class="danger" onclick="clearCopyTableCache()">清除缓存</button>
          <button type="button" onclick="loadCopyTable()">刷新</button>
        </div>
      </div>
      <div style="margin-top:10px" class="muted" id="copy-table-status"></div>
    </div>

    <div class="card">
      <table id="copy-table">
        <thead>
          <tr>
            <th>快手昵称</th>
            <th>快手ID</th>
            <th>主播</th>
            <th>直播账号</th>
            <th>添加时间</th>
            <th>操作</th>
          </tr>
        </thead>
        <tbody id="copy-tbody"></tbody>
      </table>
    </div>
  </div>

<script>
const TOKEN = new URLSearchParams(location.search).get('token') || '';

try {
  const nojs = document.getElementById('nojs');
  if (nojs) nojs.style.display = 'none';
} catch (e) {}

// Tab切换
function switchTab(tabName) {
  document.querySelectorAll('.tab').forEach(t => t.classList.remove('active'));
  document.querySelectorAll('.tab-content').forEach(t => t.classList.remove('active'));
  
  event.target.classList.add('active');
  document.getElementById('tab-' + tabName).classList.add('active');
  
  if (tabName === 'mapping') {
    loadMappingTables();
  }
  if (tabName === 'mapping_wenzong') {
    loadMappingTablesWenzong();
  }
  if (tabName === 'niu') {
    loadNiuTable();
  }
  if (tabName === 'copy') {
    loadCopyTable();
  }
}

async function loadNiuTable() {
  try {
    const data = await api('/api/niu_table');
    const rows = Array.isArray(data.rows) ? data.rows : [];
    renderNiuTable(rows);
  } catch (err) {
    alert('加载金牛表失败: ' + err.message);
  }
}

function renderNiuTable(rows) {
  const tbody = document.getElementById('niu-tbody');
  if (!tbody) return;
  tbody.innerHTML = '';
  rows.forEach((r, i) => {
    const tr = document.createElement('tr');
    const url = esc((r && r.url) ? r.url : '');
    const bg = (r && r.bg_login && typeof r.bg_login === 'object') ? r.bg_login : {};
    const running = !!bg.running;
    const ok = !!bg.ok;
    const progress = Number(bg.progress || 0);
    const msg = esc((bg.message || '').toString());
    const finishedAt = esc((bg.finished_at || bg.last_success_at || '').toString());
    const statusClass = running ? 'status-warn' : (ok ? 'status-ok' : (msg ? 'status-err' : 'muted'));
    const statusText = running
      ? `进行中 ${progress}% ${msg}`
      : (ok ? `已完成 ${finishedAt || msg}` : (msg || '未登录'));
    tr.innerHTML = `
      <td><input type="text" value="${url}" data-row="${i}" /></td>
      <td>
        <div class="${statusClass}">${statusText}</div>
      </td>
      <td>
        <button class="btn-small" onclick="openNiuUrl(${i})">打开</button>
        <button class="btn-small" onclick="openNiuLoginModal(${i})">后台登录</button>
        <button class="btn-small danger" onclick="deleteNiuRow(${i})">删除</button>
      </td>
    `;
    tbody.appendChild(tr);
  });
}

function addNiuRow() {
  const tbody = document.getElementById('niu-tbody');
  if (!tbody) return;
  const i = tbody.children.length;
  const tr = document.createElement('tr');
  tr.innerHTML = `
    <td><input type="text" value="" data-row="${i}" placeholder="https://niu.e.kuaishou.com/..." /></td>
    <td><div class="muted">未登录</div></td>
    <td>
      <button class="btn-small" onclick="openNiuUrl(${i})">打开</button>
      <button class="btn-small" onclick="openNiuLoginModal(${i})">后台登录</button>
      <button class="btn-small danger" onclick="deleteNiuRow(${i})">删除</button>
    </td>
  `;
  tbody.appendChild(tr);
}

function deleteNiuRow(i) {
  const tbody = document.getElementById('niu-tbody');
  if (!tbody) return;
  if (tbody.children[i]) tbody.children[i].remove();
  Array.from(tbody.children).forEach((tr, idx) => {
    const btns = tr.querySelectorAll('button');
    if (btns[0]) btns[0].setAttribute('onclick', `openNiuUrl(${idx})`);
    if (btns[1]) btns[1].setAttribute('onclick', `openNiuLoginModal(${idx})`);
    if (btns[2]) btns[2].setAttribute('onclick', `deleteNiuRow(${idx})`);
  });
}

async function saveNiuTable() {
  const tbody = document.getElementById('niu-tbody');
  if (!tbody) return;
  const rows = [];
  Array.from(tbody.children).forEach(tr => {
    const input = tr.querySelector('input');
    const url = input ? input.value.trim() : '';
    if (url) rows.push({ url });
  });
  try {
    await api('/api/niu_table', {
      method: 'POST',
      body: JSON.stringify({ rows })
    });
    alert('保存成功！');
    loadNiuTable();
  } catch (err) {
    alert('保存失败: ' + err.message);
  }
}

function openNiuUrl(i) {
  api(`/api/open_niu_url?idx=${encodeURIComponent(String(i))}`)
    .catch(err => {
      alert('打开失败: ' + err.message);
    });
}

const niuLoginState = { idx: -1, timer: null };
const kuaishouLoginState = { account: '', ksId: '', scope: '', timer: null };

function setNiuLoginProgress(progress, text) {
  const bar = document.getElementById('niu-login-progress-bar');
  const textEl = document.getElementById('niu-login-progress-text');
  if (bar) bar.style.width = `${Math.max(0, Math.min(100, Number(progress || 0)))}%`;
  if (textEl) textEl.textContent = text || '';
}

function closeNiuLoginModal() {
  const modal = document.getElementById('niu-login-modal');
  if (modal) modal.classList.remove('show');
  if (niuLoginState.timer) {
    clearInterval(niuLoginState.timer);
    niuLoginState.timer = null;
  }
  niuLoginState.idx = -1;
}

function openNiuLoginModal(i) {
  const tbody = document.getElementById('niu-tbody');
  const modal = document.getElementById('niu-login-modal');
  const urlEl = document.getElementById('niu-login-modal-url');
  const row = tbody && tbody.children[i] ? tbody.children[i] : null;
  const input = row ? row.querySelector('input[type="text"]') : null;
  const url = input ? input.value.trim() : '';
  niuLoginState.idx = i;
  if (urlEl) urlEl.textContent = url ? `当前网址: ${url}` : '请先填写并保存金牛网址';
  setNiuLoginProgress(0, '等待开始');
  if (modal) modal.classList.add('show');
}

async function pollNiuLoginStatus() {
  if (niuLoginState.idx < 0) return;
  try {
    const data = await api(`/api/niu_login_status?idx=${encodeURIComponent(String(niuLoginState.idx))}`);
    const st = data.status || {};
    setNiuLoginProgress(st.progress || 0, st.message || '');
    if (st.running) return;
    if (niuLoginState.timer) {
      clearInterval(niuLoginState.timer);
      niuLoginState.timer = null;
    }
    if (st.ok) {
      setNiuLoginProgress(100, st.message || '后台登录完成');
      setTimeout(() => {
        closeNiuLoginModal();
        loadNiuTable();
      }, 800);
      return;
    }
    loadNiuTable();
  } catch (err) {
    if (niuLoginState.timer) {
      clearInterval(niuLoginState.timer);
      niuLoginState.timer = null;
    }
    setNiuLoginProgress(100, '后台登录失败: ' + err.message);
  }
}

async function submitNiuBackgroundLogin() {
  if (niuLoginState.idx < 0) return;
  const phoneEl = document.getElementById('niu-login-phone');
  const passwordEl = document.getElementById('niu-login-password');
  const phone = phoneEl ? phoneEl.value.trim() : '';
  const password = passwordEl ? passwordEl.value : '';
  if (!phone || !password) {
    alert('请输入账号和密码');
    return;
  }
  setNiuLoginProgress(5, '正在创建后台登录任务');
  try {
    const data = await api('/api/niu_login_background', {
      method: 'POST',
      body: JSON.stringify({ idx: niuLoginState.idx, phone, password }),
    });
    const st = data.status || {};
    setNiuLoginProgress(st.progress || 5, st.message || '后台登录已启动');
    if (niuLoginState.timer) clearInterval(niuLoginState.timer);
    niuLoginState.timer = setInterval(pollNiuLoginStatus, 1000);
    pollNiuLoginStatus();
  } catch (err) {
    setNiuLoginProgress(100, '后台登录失败: ' + err.message);
  }
}

function setKuaishouLoginProgress(progress, text) {
  const bar = document.getElementById('kuaishou-login-progress-bar');
  const textEl = document.getElementById('kuaishou-login-progress-text');
  if (bar) bar.style.width = `${Math.max(0, Math.min(100, Number(progress || 0)))}%`;
  if (textEl) textEl.textContent = text || '';
}

function setKuaishouQr(url) {
  const img = document.getElementById('kuaishou-login-qr');
  const empty = document.getElementById('kuaishou-login-qr-empty');
  if (!img || !empty) return;
  if (url) {
    img.src = url;
    img.style.display = 'block';
    empty.style.display = 'none';
  } else {
    img.removeAttribute('src');
    img.style.display = 'none';
    empty.style.display = 'block';
  }
}

function closeKuaishouLoginModal() {
  const modal = document.getElementById('kuaishou-login-modal');
  if (modal) modal.classList.remove('show');
  if (kuaishouLoginState.timer) {
    clearInterval(kuaishouLoginState.timer);
    kuaishouLoginState.timer = null;
  }
  kuaishouLoginState.account = '';
  kuaishouLoginState.ksId = '';
  kuaishouLoginState.scope = '';
  setKuaishouQr('');
}

function openKuaishouLoginModal(accountName, ksId, scope='') {
  if (!accountName || !ksId) {
    alert('账号信息不完整');
    return;
  }
  kuaishouLoginState.account = accountName;
  kuaishouLoginState.ksId = ksId;
  kuaishouLoginState.scope = scope || '';
  const modal = document.getElementById('kuaishou-login-modal');
  const titleEl = document.getElementById('kuaishou-login-modal-title');
  if (titleEl) {
    titleEl.textContent = `直播账号: ${accountName} | 快手ID: ${ksId}`;
  }
  setKuaishouLoginProgress(0, '等待开始');
  setKuaishouQr('');
  if (modal) modal.classList.add('show');
}

async function pollKuaishouLoginStatus() {
  if (!kuaishouLoginState.ksId) return;
  const scope = kuaishouLoginState.scope || '';
  try {
    const data = await api(`/api/kuaishou_login_status?scope=${encodeURIComponent(scope)}&ks_id=${encodeURIComponent(kuaishouLoginState.ksId)}`);
    const st = data.status || {};
    setKuaishouLoginProgress(st.progress || 0, st.message || '');
    if (st.qr_ready) {
      setKuaishouQr(`/api/kuaishou_login_qr?scope=${encodeURIComponent(scope)}&ks_id=${encodeURIComponent(kuaishouLoginState.ksId)}&ts=${Date.now()}${TOKEN ? ('&token=' + encodeURIComponent(TOKEN)) : ''}`);
    } else if (st.ok) {
      setKuaishouQr('');
    }
    if (st.running) return;
    if (kuaishouLoginState.timer) {
      clearInterval(kuaishouLoginState.timer);
      kuaishouLoginState.timer = null;
    }
    if (st.ok) {
      setKuaishouLoginProgress(100, st.message || '后台登录完成');
      setTimeout(() => {
        closeKuaishouLoginModal();
      }, 1000);
    }
  } catch (err) {
    if (kuaishouLoginState.timer) {
      clearInterval(kuaishouLoginState.timer);
      kuaishouLoginState.timer = null;
    }
    setKuaishouLoginProgress(100, '后台登录失败: ' + err.message);
  }
}

async function submitKuaishouBackgroundLogin() {
  if (!kuaishouLoginState.account || !kuaishouLoginState.ksId) return;
  setKuaishouLoginProgress(5, '正在创建后台登录任务');
  try {
    const data = await api('/api/kuaishou_login_background', {
      method: 'POST',
      body: JSON.stringify({
        scope: kuaishouLoginState.scope || '',
        account: kuaishouLoginState.account,
        ks_id: kuaishouLoginState.ksId,
      }),
    });
    const st = data.status || {};
    setKuaishouLoginProgress(st.progress || 5, st.message || '后台登录已启动');
    if (st.qr_ready) {
      setKuaishouQr(`/api/kuaishou_login_qr?scope=${encodeURIComponent(kuaishouLoginState.scope || '')}&ks_id=${encodeURIComponent(kuaishouLoginState.ksId)}&ts=${Date.now()}${TOKEN ? ('&token=' + encodeURIComponent(TOKEN)) : ''}`);
    }
    if (kuaishouLoginState.timer) clearInterval(kuaishouLoginState.timer);
    kuaishouLoginState.timer = setInterval(pollKuaishouLoginStatus, 1500);
    pollKuaishouLoginStatus();
  } catch (err) {
    setKuaishouLoginProgress(100, '后台登录失败: ' + err.message);
  }
}

async function loadWenzongConfig() {
  try {
    const data = await api('/api/wenzong/config');
    const enabled = !!(data && data.enabled);
    const elEnabled = document.getElementById('wenzong-enabled');
    if (elEnabled) elEnabled.checked = enabled;
  } catch (err) {
    // ignore
  }
}

async function saveWenzongEnabledOnly() {
  const elEnabled = document.getElementById('wenzong-enabled');
  const enabled = !!(elEnabled && elEnabled.checked);
  try {
    await api('/api/wenzong/config', {
      method: 'POST',
      body: JSON.stringify({ enabled })
    });
  } catch (err) {
    alert('保存失败: ' + err.message);
  }
}

async function loadCopyTable() {
  const statusEl = document.getElementById('copy-table-status');
  try {
    const data = await api('/api/copy_table');
    const items = Array.isArray(data.items) ? data.items : [];
    items.sort((a, b) => {
      const ta = (a && a.created_at) ? String(a.created_at) : '';
      const tb = (b && b.created_at) ? String(b.created_at) : '';
      if (ta && tb) return tb.localeCompare(ta);
      if (!ta && tb) return 1;
      if (ta && !tb) return -1;
      return 0;
    });
    renderCopyTable(items);
    if (statusEl) statusEl.textContent = `共 ${items.length} 条 | updated_at: ${data.updated_at||''}`;
  } catch (err) {
    if (statusEl) statusEl.textContent = '加载失败: ' + err.message;
  }
}

function renderCopyTable(items) {
  const tbody = document.getElementById('copy-tbody');
  if (!tbody) return;
  tbody.innerHTML = '';
  items.forEach(it => {
    const tr = document.createElement('tr');
    const nick = (it && it.nickname) ? String(it.nickname) : '';
    const kid = (it && it.kuaishou_id) ? String(it.kuaishou_id) : '';
    const anchor = (it && it.anchor) ? String(it.anchor) : '';
    const acct = (it && it.account) ? String(it.account) : '';
    const createdAt = (it && it.created_at) ? String(it.created_at) : '';
    const copied = !!(it && it.copied);
    const key = (it && it.key) ? String(it.key) : '';

    const td1 = document.createElement('td'); td1.textContent = nick;
    const td2 = document.createElement('td'); td2.textContent = kid;
    const td3 = document.createElement('td'); td3.textContent = anchor;
    const td4 = document.createElement('td'); td4.textContent = acct;
    const td6 = document.createElement('td'); td6.textContent = createdAt;
    const td5 = document.createElement('td');

    const btn = document.createElement('button');
    btn.type = 'button';
    btn.className = 'btn-small' + (copied ? '' : '');
    btn.textContent = copied ? '复制（已复制）' : '复制';
    btn.onclick = async () => {
      const text = `快手昵称：${nick}\n快手ID：${kid}\n主播：${anchor}\n直播账号：${acct}`;
      await copyTextToClipboard(text);
      try {
        await api('/api/copy_table/mark_copied', { method: 'POST', body: JSON.stringify({ key, copied: true }) });
      } catch (e) {}
      loadCopyTable();
    };
    td5.appendChild(btn);

    tr.appendChild(td1); tr.appendChild(td2); tr.appendChild(td3); tr.appendChild(td4); tr.appendChild(td6); tr.appendChild(td5);
    tbody.appendChild(tr);
  });
}

async function copyTextToClipboard(text) {
  text = text || '';
  if (navigator.clipboard && navigator.clipboard.writeText) {
    return navigator.clipboard.writeText(text);
  }
  const ta = document.createElement('textarea');
  ta.value = text;
  ta.style.position = 'fixed';
  ta.style.left = '-9999px';
  document.body.appendChild(ta);
  ta.focus();
  ta.select();
  try {
    document.execCommand('copy');
  } finally {
    document.body.removeChild(ta);
  }
}

async function clearCopyTableCache() {
  if (!confirm('确定清除所有复制数据缓存吗？')) return;
  try {
    await api('/api/copy_table/clear', { method: 'POST', body: JSON.stringify({}) });
    loadCopyTable();
  } catch (err) {
    alert('清除失败: ' + err.message);
  }
}

async function loadUserContactDept() {
  try {
    const data = await api('/api/user_contact_config');
    const el = document.getElementById('uc-dept');
    if (el) el.value = (data.dept || '');
  } catch (err) {
    // ignore
  }
}

async function loadWenzongUserContactDept() {
  try {
    const data = await api('/api/user_contact_config?scope=wenzong');
    const el = document.getElementById('wz-uc-dept');
    if (el) el.value = (data.dept || '');
  } catch (err) {
    // ignore
  }
}

async function saveUserContactDept() {
  const el = document.getElementById('uc-dept');
  const dept = (el && el.value) ? el.value.trim() : '';
  try {
    await api('/api/user_contact_config', {
      method: 'POST',
      body: JSON.stringify({ dept })
    });
    alert('部门已保存');
  } catch (err) {
    alert('保存失败: ' + err.message);
  }
}

async function saveWenzongUserContactDept() {
  const el = document.getElementById('wz-uc-dept');
  const dept = (el && el.value) ? el.value.trim() : '';
  try {
    await api('/api/user_contact_config?scope=wenzong', {
      method: 'POST',
      body: JSON.stringify({ dept })
    });
    alert('WenZong部门已保存');
  } catch (err) {
    alert('保存失败: ' + err.message);
  }
}

async function importUserContactXlsx(ev) {
  const file = ev && ev.target && ev.target.files && ev.target.files[0];
  if (!file) return;
  const fd = new FormData();
  fd.append('file', file);
  try {
    const resp = await fetch('/api/user_contact/import_xlsx', {
      method: 'POST',
      headers: { 'X-Token': TOKEN },
      body: fd
    });
    const data = await resp.json().catch(() => ({}));
    if (!resp.ok || !data.ok) throw new Error(data.error || ('http_' + resp.status));
    alert('导入成功');
  } catch (err) {
    alert('导入失败: ' + err.message);
  } finally {
    try { ev.target.value = ''; } catch (e) {}
  }
}

function exportUserContactXlsx() {
  const url = '/api/user_contact/export_xlsx' + (TOKEN ? ('?token=' + encodeURIComponent(TOKEN)) : '');
  window.open(url, '_blank');
}

// 加载主播映射表（两个表）
async function loadMappingTables() {
  try {
    const data = await api('/api/mapping');
    renderMappingTable('active', data.active_rows || []);
    renderMappingTable('inactive', data.inactive_rows || []);
  } catch (err) {
    alert('加载映射表失败: ' + err.message);
  }
}

async function loadMappingTablesWenzong() {
  try {
    const data = await api('/api/mapping?scope=wenzong');
    renderMappingTableWenzong('active', data.active_rows || []);
    renderMappingTableWenzong('inactive', data.inactive_rows || []);
  } catch (err) {
    alert('加载 WenZong 映射表失败: ' + err.message);
  }
}

function renderMappingTableWenzong(tableType, rows) {
  const tbody = document.getElementById('mapping-tbody-' + tableType + '-wenzong');
  tbody.innerHTML = '';
  
  rows.forEach((row, index) => {
    const tr = document.createElement('tr');
    const otherType = tableType === 'active' ? 'inactive' : 'active';
    const moveLabel = tableType === 'active' ? '移至暂不使用' : '恢复使用';
    const accountName = esc(row[0] || '');
    const ksId = esc(row[1] || '');
    
    tr.innerHTML = `
      <td><input type="text" value="${accountName}" data-col="0" data-row="${index}" /></td>
      <td><input type="text" value="${ksId}" data-col="1" data-row="${index}" /></td>
      <td><input type="text" value="${esc(row[2] || '')}" data-col="2" data-row="${index}" /></td>
      <td><input type="text" value="${esc(row[3] || '')}" data-col="3" data-row="${index}" /></td>
      <td><input type="text" value="${esc(row[4] || '')}" data-col="4" data-row="${index}" /></td>
      <td>
        <button class="btn-small" onclick="openKuaishouClassroomWenzong('${accountName}', '${ksId}')">打开课堂</button>
        <button class="btn-small" onclick="openKuaishouLoginModal('${accountName}', '${ksId}', 'wenzong')">后台登录</button>
      </td>
      <td>
        <button class="btn-small" onclick="moveMappingRowWenzong('${tableType}', ${index}, '${otherType}')">${moveLabel}</button>
        <button class="btn-small danger" onclick="deleteMappingRowWenzong('${tableType}', ${index})">删除</button>
      </td>
    `;
    tbody.appendChild(tr);
  });
}

function addMappingRowWenzong(tableType) {
  const tbody = document.getElementById('mapping-tbody-' + tableType + '-wenzong');
  const index = tbody.children.length;
  const otherType = tableType === 'active' ? 'inactive' : 'active';
  const moveLabel = tableType === 'active' ? '移至暂不使用' : '恢复使用';
  
  const tr = document.createElement('tr');
  tr.innerHTML = `
    <td><input type="text" value="" data-col="0" data-row="${index}" placeholder="直播账号" /></td>
    <td><input type="text" value="" data-col="1" data-row="${index}" placeholder="快手ID" /></td>
    <td><input type="text" value="" data-col="2" data-row="${index}" placeholder="手机号码" /></td>
    <td><input type="text" value="" data-col="3" data-row="${index}" placeholder="密码" /></td>
    <td><input type="text" value="" data-col="4" data-row="${index}" placeholder="主播" /></td>
    <td>
      <button class="btn-small" onclick="openKuaishouClassroomWenzong('', '')">打开课堂</button>
      <button class="btn-small" onclick="openKuaishouLoginModal('', '', 'wenzong')">后台登录</button>
    </td>
    <td>
      <button class="btn-small" onclick="moveMappingRowWenzong('${tableType}', ${index}, '${otherType}')">${moveLabel}</button>
      <button class="btn-small danger" onclick="deleteMappingRowWenzong('${tableType}', ${index})">删除</button>
    </td>
  `;
  tbody.appendChild(tr);
}

function openKuaishouClassroomWenzong(accountName, ksId) {
  if (!accountName || !ksId) {
    alert('账号信息不完整');
    return;
  }
  api(`/api/open_kuaishou_classroom?scope=wenzong&account=${encodeURIComponent(accountName)}&ks_id=${encodeURIComponent(ksId)}`)
    .then(data => {
      console.log(`打开快手课堂(WenZong) - 账号: ${accountName}, 快手ID: ${ksId}`);
      const statusEl = document.getElementById('cache-status');
      if (statusEl) {
        statusEl.textContent = `正在打开 ${accountName} 的快手课堂...`;
        setTimeout(() => {
          statusEl.textContent = '';
        }, 3000);
      }
    })
    .catch(err => {
      alert('打开快手课堂失败: ' + err.message);
      console.error('Error opening classroom:', err);
    });
}

function deleteMappingRowWenzong(tableType, index) {
  const tbody = document.getElementById('mapping-tbody-' + tableType + '-wenzong');
  if (tbody.children[index]) {
    tbody.children[index].remove();
    updateRowIndicesWenzong(tableType);
  }
}

function moveMappingRowWenzong(fromType, index, toType) {
  const fromTbody = document.getElementById('mapping-tbody-' + fromType + '-wenzong');
  const toTbody = document.getElementById('mapping-tbody-' + toType + '-wenzong');
  if (!fromTbody.children[index]) return;
  
  const tr = fromTbody.children[index];
  const inputs = tr.querySelectorAll('input');
  const rowData = Array.from(inputs).map(input => input.value.trim());
  
  tr.remove();
  updateRowIndicesWenzong(fromType);
  
  const newIndex = toTbody.children.length;
  const newTr = document.createElement('tr');
  const moveLabel = toType === 'active' ? '移至暂不使用' : '恢复使用';
  const otherType = toType === 'active' ? 'inactive' : 'active';
  const accountName = esc(rowData[0] || '');
  const ksId = esc(rowData[1] || '');
  
  newTr.innerHTML = `
    <td><input type="text" value="${accountName}" data-col="0" data-row="${newIndex}" /></td>
    <td><input type="text" value="${ksId}" data-col="1" data-row="${newIndex}" /></td>
    <td><input type="text" value="${esc(rowData[2] || '')}" data-col="2" data-row="${newIndex}" /></td>
    <td><input type="text" value="${esc(rowData[3] || '')}" data-col="3" data-row="${newIndex}" /></td>
    <td><input type="text" value="${esc(rowData[4] || '')}" data-col="4" data-row="${newIndex}" /></td>
    <td>
      <button class="btn-small" onclick="openKuaishouClassroomWenzong('${accountName}', '${ksId}')">打开课堂</button>
      <button class="btn-small" onclick="openKuaishouLoginModal('${accountName}', '${ksId}', 'wenzong')">后台登录</button>
    </td>
    <td>
      <button class="btn-small" onclick="moveMappingRowWenzong('${toType}', ${newIndex}, '${otherType}')">${moveLabel}</button>
      <button class="btn-small danger" onclick="deleteMappingRowWenzong('${toType}', ${newIndex})">删除</button>
    </td>
  `;
  toTbody.appendChild(newTr);
}

function updateRowIndicesWenzong(tableType) {
  const tbody = document.getElementById('mapping-tbody-' + tableType + '-wenzong');
  const otherType = tableType === 'active' ? 'inactive' : 'active';
  
  Array.from(tbody.children).forEach((tr, newIndex) => {
    tr.querySelectorAll('input').forEach(input => {
      input.setAttribute('data-row', newIndex);
    });
    
    const inputs = tr.querySelectorAll('input');
    const accountName = inputs[0] ? inputs[0].value.trim() : '';
    const ksId = inputs[1] ? inputs[1].value.trim() : '';
    
    const buttons = tr.querySelectorAll('button');
    if (buttons[0]) buttons[0].setAttribute('onclick', `openKuaishouClassroomWenzong('${esc(accountName)}', '${esc(ksId)}')`);
    if (buttons[1]) buttons[1].setAttribute('onclick', `openKuaishouLoginModal('${esc(accountName)}', '${esc(ksId)}', 'wenzong')`);
    if (buttons[2]) buttons[2].setAttribute('onclick', `moveMappingRowWenzong('${tableType}', ${newIndex}, '${otherType}')`);
    if (buttons[3]) buttons[3].setAttribute('onclick', `deleteMappingRowWenzong('${tableType}', ${newIndex})`);
  });
}

async function saveMappingTablesWenzong() {
  const activeRows = [];
  const inactiveRows = [];
  
  const activeTbody = document.getElementById('mapping-tbody-active-wenzong');
  Array.from(activeTbody.children).forEach(tr => {
    const inputs = tr.querySelectorAll('input');
    const row = Array.from(inputs).map(input => input.value.trim());
    if (row.some(v => v)) {
      activeRows.push(row);
    }
  });
  
  const inactiveTbody = document.getElementById('mapping-tbody-inactive-wenzong');
  Array.from(inactiveTbody.children).forEach(tr => {
    const inputs = tr.querySelectorAll('input');
    const row = Array.from(inputs).map(input => input.value.trim());
    if (row.some(v => v)) {
      inactiveRows.push(row);
    }
  });
  
  try {
    await api('/api/mapping?scope=wenzong', {
      method: 'POST',
      body: JSON.stringify({ 
        active_rows: activeRows,
        inactive_rows: inactiveRows
      })
    });
    alert('保存成功！');
    loadMappingTablesWenzong();
    refresh();
  } catch (err) {
    alert('保存失败: ' + err.message);
  }
}

// 渲染映射表
function renderMappingTable(tableType, rows) {
  const tbody = document.getElementById('mapping-tbody-' + tableType);
  tbody.innerHTML = '';
  
  rows.forEach((row, index) => {
    const tr = document.createElement('tr');
    const otherType = tableType === 'active' ? 'inactive' : 'active';
    const moveLabel = tableType === 'active' ? '移至暂不使用' : '恢复使用';
    const accountName = esc(row[0] || '');
    const ksId = esc(row[1] || '');
    
    tr.innerHTML = `
      <td><input type="text" value="${accountName}" data-col="0" data-row="${index}" /></td>
      <td><input type="text" value="${ksId}" data-col="1" data-row="${index}" /></td>
      <td><input type="text" value="${esc(row[2] || '')}" data-col="2" data-row="${index}" /></td>
      <td><input type="text" value="${esc(row[3] || '')}" data-col="3" data-row="${index}" /></td>
      <td><input type="text" value="${esc(row[4] || '')}" data-col="4" data-row="${index}" /></td>
      <td>
        <button class="btn-small" onclick="openKuaishouClassroom('${accountName}', '${ksId}')">打开课堂</button>
        <button class="btn-small" onclick="openKuaishouLoginModal('${accountName}', '${ksId}')">后台登录</button>
      </td>
      <td>
        <button class="btn-small" onclick="moveMappingRow('${tableType}', ${index}, '${otherType}')">${moveLabel}</button>
        <button class="btn-small danger" onclick="deleteMappingRow('${tableType}', ${index})">删除</button>
      </td>
    `;
    tbody.appendChild(tr);
  });
}

// 新增行
function addMappingRow(tableType) {
  const tbody = document.getElementById('mapping-tbody-' + tableType);
  const index = tbody.children.length;
  const otherType = tableType === 'active' ? 'inactive' : 'active';
  const moveLabel = tableType === 'active' ? '移至暂不使用' : '恢复使用';
  
  const tr = document.createElement('tr');
  tr.innerHTML = `
    <td><input type="text" value="" data-col="0" data-row="${index}" placeholder="直播账号" /></td>
    <td><input type="text" value="" data-col="1" data-row="${index}" placeholder="快手ID" /></td>
    <td><input type="text" value="" data-col="2" data-row="${index}" placeholder="手机号码" /></td>
    <td><input type="text" value="" data-col="3" data-row="${index}" placeholder="密码" /></td>
    <td><input type="text" value="" data-col="4" data-row="${index}" placeholder="主播" /></td>
    <td>
      <button class="btn-small" onclick="openKuaishouClassroom('', '')">打开课堂</button>
      <button class="btn-small" onclick="openKuaishouLoginModal('', '')">后台登录</button>
    </td>
    <td>
      <button class="btn-small" onclick="moveMappingRow('${tableType}', ${index}, '${otherType}')">${moveLabel}</button>
      <button class="btn-small danger" onclick="deleteMappingRow('${tableType}', ${index})">删除</button>
    </td>
  `;
  tbody.appendChild(tr);
}

// 打开快手课堂
function openKuaishouClassroom(accountName, ksId) {
  if (!accountName || !ksId) {
    alert('账号信息不完整');
    return;
  }
  
  // 调用API打开带登录状态的浏览器
  api(`/api/open_kuaishou_classroom?account=${encodeURIComponent(accountName)}&ks_id=${encodeURIComponent(ksId)}`)
    .then(data => {
      console.log(`打开快手课堂 - 账号: ${accountName}, 快手ID: ${ksId}`);
      // 显示提示信息
      const statusEl = document.getElementById('cache-status');
      if (statusEl) {
        statusEl.textContent = `正在打开 ${accountName} 的快手课堂...`;
        setTimeout(() => {
          statusEl.textContent = '';
        }, 3000);
      }
    })
    .catch(err => {
      alert('打开快手课堂失败: ' + err.message);
      console.error('Error opening classroom:', err);
    });
}

// 删除行
function deleteMappingRow(tableType, index) {
  const tbody = document.getElementById('mapping-tbody-' + tableType);
  if (tbody.children[index]) {
    tbody.children[index].remove();
    // 重新编号
    updateRowIndices(tableType);
  }
}

// 移动行到另一个表
function moveMappingRow(fromType, index, toType) {
  const fromTbody = document.getElementById('mapping-tbody-' + fromType);
  const toTbody = document.getElementById('mapping-tbody-' + toType);
  
  if (!fromTbody.children[index]) return;
  
  const tr = fromTbody.children[index];
  const inputs = tr.querySelectorAll('input');
  const rowData = Array.from(inputs).map(input => input.value.trim());
  
  // 从源表删除
  tr.remove();
  updateRowIndices(fromType);
  
  // 添加到目标表
  const newIndex = toTbody.children.length;
  const newTr = document.createElement('tr');
  const moveLabel = toType === 'active' ? '移至暂不使用' : '恢复使用';
  const otherType = toType === 'active' ? 'inactive' : 'active';
  const accountName = esc(rowData[0] || '');
  const ksId = esc(rowData[1] || '');
  
  newTr.innerHTML = `
    <td><input type="text" value="${accountName}" data-col="0" data-row="${newIndex}" /></td>
    <td><input type="text" value="${ksId}" data-col="1" data-row="${newIndex}" /></td>
    <td><input type="text" value="${esc(rowData[2] || '')}" data-col="2" data-row="${newIndex}" /></td>
    <td><input type="text" value="${esc(rowData[3] || '')}" data-col="3" data-row="${newIndex}" /></td>
    <td><input type="text" value="${esc(rowData[4] || '')}" data-col="4" data-row="${newIndex}" /></td>
    <td>
      <button class="btn-small" onclick="openKuaishouClassroom('${accountName}', '${ksId}')">打开课堂</button>
      <button class="btn-small" onclick="openKuaishouLoginModal('${accountName}', '${ksId}')">后台登录</button>
    </td>
    <td>
      <button class="btn-small" onclick="moveMappingRow('${toType}', ${newIndex}, '${otherType}')">${moveLabel}</button>
      <button class="btn-small danger" onclick="deleteMappingRow('${toType}', ${newIndex})">删除</button>
    </td>
  `;
  toTbody.appendChild(newTr);
}

// 更新行索引
function updateRowIndices(tableType) {
  const tbody = document.getElementById('mapping-tbody-' + tableType);
  const otherType = tableType === 'active' ? 'inactive' : 'active';
  const moveLabel = tableType === 'active' ? '移至暂不使用' : '恢复使用';
  
  Array.from(tbody.children).forEach((tr, newIndex) => {
    tr.querySelectorAll('input').forEach(input => {
      input.setAttribute('data-row', newIndex);
    });
    
    // 获取账号名和快手ID用于快手课堂按钮
    const inputs = tr.querySelectorAll('input');
    const accountName = inputs[0] ? inputs[0].value.trim() : '';
    const ksId = inputs[1] ? inputs[1].value.trim() : '';
    
    const buttons = tr.querySelectorAll('button');
    if (buttons[0]) buttons[0].setAttribute('onclick', `openKuaishouClassroom('${esc(accountName)}', '${esc(ksId)}')`);
    if (buttons[1]) buttons[1].setAttribute('onclick', `openKuaishouLoginModal('${esc(accountName)}', '${esc(ksId)}')`);
    if (buttons[2]) buttons[2].setAttribute('onclick', `moveMappingRow('${tableType}', ${newIndex}, '${otherType}')`);
    if (buttons[3]) buttons[3].setAttribute('onclick', `deleteMappingRow('${tableType}', ${newIndex})`);
  });
}

// 保存映射表（两个表）
async function saveMappingTables() {
  const activeRows = [];
  const inactiveRows = [];
  
  // 收集活跃表数据
  const activeTbody = document.getElementById('mapping-tbody-active');
  Array.from(activeTbody.children).forEach(tr => {
    const inputs = tr.querySelectorAll('input');
    const row = Array.from(inputs).map(input => input.value.trim());
    if (row.some(v => v)) {
      activeRows.push(row);
    }
  });
  
  // 收集暂不使用表数据
  const inactiveTbody = document.getElementById('mapping-tbody-inactive');
  Array.from(inactiveTbody.children).forEach(tr => {
    const inputs = tr.querySelectorAll('input');
    const row = Array.from(inputs).map(input => input.value.trim());
    if (row.some(v => v)) {
      inactiveRows.push(row);
    }
  });
  
  try {
    const result = await api('/api/mapping', {
      method: 'POST',
      body: JSON.stringify({ 
        active_rows: activeRows,
        inactive_rows: inactiveRows
      })
    });
    alert('保存成功！');
    loadMappingTables();
    refresh();
  } catch (err) {
    alert('保存失败: ' + err.message);
  }
}

async function api(path, opts={}) {
  const url = path + (path.includes('?') ? '&' : '?') + 'token=' + encodeURIComponent(TOKEN);
  const resp = await fetch(url, { ...opts, headers: { 'Content-Type': 'application/json', ...(opts.headers||{}) } });
  const data = await resp.json().catch(() => ({}));
  if (!resp.ok) {
    throw new Error((data && (data.error||data.msg)) || ('HTTP ' + resp.status));
  }
  return data;
}

function esc(s) {
  return (s||'').toString().replace(/[&<>"']/g, c => {
    const map = {'&':'&amp;', '<':'&lt;', '>':'&gt;', '"':'&quot;', "'":"&#39;"};
    return map[c];
  });
}

function renderNiuRuntimeStatus(items) {
  const el = document.getElementById('niu-runtime-status');
  if (!el) return;
  const arr = Array.isArray(items) ? items : [];
  if (!arr.length) {
    el.innerHTML = '<div class="muted">暂无金牛运行状态</div>';
    return;
  }
  el.innerHTML = arr.map(item => {
    const label = esc(item.label || item.key || '金牛');
    const msg = esc(item.message || '等待运行');
    const progress = Number(item.progress || 0);
    const updated = esc(item.updated_at || '');
    const cls = item.running ? 'status-warn' : (item.ok ? 'status-ok' : 'status-err');
    return `<div class="${cls}" style="margin-bottom:6px">${label}: ${msg} (${progress}%)${updated ? ' | ' + updated : ''}</div>`;
  }).join('');
}

function renderKuaishouRuntimeStatus(items) {
  const el = document.getElementById('kuaishou-runtime-status');
  if (!el) return;
  const arr = Array.isArray(items) ? items : [];
  if (!arr.length) {
    el.innerHTML = '<div class="muted">暂无快手课堂运行状态</div>';
    return;
  }
  el.innerHTML = arr.map(item => {
    const label = esc(item.label || item.key || '快手课堂');
    const msg = esc(item.message || '等待运行');
    const progress = Number(item.progress || 0);
    const updated = esc(item.updated_at || '');
    const cls = item.running ? 'status-warn' : (item.ok ? 'status-ok' : 'status-err');
    return `<div class="${cls}" style="margin-bottom:6px">${label}: ${msg} (${progress}%)${updated ? ' | ' + updated : ''}</div>`;
  }).join('');
}

function renderFeishuRuntimeStatus(items) {
  const el = document.getElementById('feishu-runtime-status');
  if (!el) return;
  const arr = Array.isArray(items) ? items : [];
  if (!arr.length) {
    el.innerHTML = '<div class="muted">暂无云文档与导出运行状态</div>';
    return;
  }
  arr.sort((a, b) => {
    const order = { summary: 0, user_contact: 1, xlsx_export: 2, delivery: 3 };
    const ak = String(a.key || '');
    const bk = String(b.key || '');
    return (order[ak] ?? 99) - (order[bk] ?? 99);
  });
  el.innerHTML = arr.map(item => {
    const label = esc(item.label || item.key || '云文档任务');
    const msg = esc(item.message || '等待运行');
    const progress = Number(item.progress || 0);
    const updated = esc(item.updated_at || '');
    const total = Number(item.total_steps || 0);
    const done = Number(item.completed_steps || 0);
    const stepText = total > 0 ? ` | 步骤 ${done}/${total}` : '';
    const cls = item.running ? 'status-warn' : (item.ok ? 'status-ok' : 'status-err');
    return `<div class="${cls}" style="margin-bottom:6px">${label}: ${msg} (${progress}%)${stepText}${updated ? ' | ' + updated : ''}</div>`;
  }).join('');
}

async function refresh() {
  const data = await api('/api/global/status');
  const s = data.status || {};
  const running = !!s.running;
  const interval = data.interval_seconds || 120;
  const nextRun = (typeof data.next_run_in_seconds === 'number') ? data.next_run_in_seconds : null;
  
  const statusEl = document.getElementById('global-status');
  const errEl = document.getElementById('global-error');
  const intervalInput = document.getElementById('interval-input');
  const nextEl = document.getElementById('next-run');
  
  if (statusEl) {
    statusEl.textContent = `状态: ${running ? '运行中' : '已停止'} | 间隔: ${interval}秒 | last_started: ${s.last_started_at||''} | last_finished: ${s.last_finished_at||''} | rc: ${s.last_returncode===null?'':s.last_returncode} | log: ${s.last_log_path||''}`;
  }
  if (nextEl) {
    window.__nextRunRemain = nextRun;
    if (!running || nextRun === null) {
      nextEl.textContent = '';
    } else {
      nextEl.textContent = `距离下一次运行: ${nextRun} 秒`;
    }
  }
  if (errEl) {
    errEl.innerHTML = s.last_error ? ('<pre>' + esc(s.last_error) + '</pre>') : '';
  }
  renderNiuRuntimeStatus(data.niu_runtime_items || []);
  renderKuaishouRuntimeStatus(data.kuaishou_runtime_items || []);
  renderFeishuRuntimeStatus(data.feishu_runtime_items || []);
  if (intervalInput) {
    intervalInput.value = interval;
  }
  const nightSleepToggle = document.getElementById('night-sleep-toggle');
  const nightSleepKnob = document.getElementById('night-sleep-knob');
  const nightSleepBg = nightSleepToggle ? nightSleepToggle.nextElementSibling : null;
  const nightSleepStatus = document.getElementById('night-sleep-status');
  const nsEnabled = !!data.night_sleep_enabled;
  const nsSleeping = !!data.sleeping;
  if (nightSleepToggle) {
    nightSleepToggle.checked = nsEnabled;
  }
  if (nightSleepKnob) {
    nightSleepKnob.style.left = nsEnabled ? '25px' : '3px';
  }
  if (nightSleepBg) {
    nightSleepBg.style.background = nsEnabled ? '#0f2a1c' : '#2a1010';
    nightSleepBg.style.borderColor = nsEnabled ? '#1d6b3d' : '#6b1b1b';
  }
  if (nightSleepStatus) {
    if (!nsEnabled) {
      nightSleepStatus.textContent = '';
    } else if (nsSleeping) {
      nightSleepStatus.textContent = '💤 当前正在休眠中，10:00 后自动恢复';
    } else {
      nightSleepStatus.textContent = '✅ 已启用，将在 04:00~10:00 暂停';
    }
  }
}

async function toggleNightSleep() {
  const cb = document.getElementById('night-sleep-toggle');
  const enabled = cb ? cb.checked : false;
  try {
    await api('/api/global/night_sleep', {
      method: 'POST',
      body: JSON.stringify({ enabled: enabled })
    });
    refresh();
  } catch (err) {
    alert('设置失败: ' + err.message);
    if (cb) cb.checked = !enabled;
    refresh();
  }
}

async function saveInterval() {
  const intervalInput = document.getElementById('interval-input');
  const seconds = parseInt(intervalInput.value) || 120;
  
  if (seconds < 1 || seconds > 3600) {
    alert('间隔时间必须在1-3600秒之间');
    return;
  }
  
  try {
    await api('/api/global/set_interval', {
      method: 'POST',
      body: JSON.stringify({ interval_seconds: seconds })
    });
    alert(`间隔时间已设置为 ${seconds} 秒`);
    refresh();
  } catch (err) {
    alert('设置失败: ' + err.message);
  }
}

async function startGlobal() {
  await api('/api/global/start', { method: 'POST', body: JSON.stringify({}) });
  setTimeout(() => { refresh().catch(() => {}); }, 300);
}

async function stopGlobal() {
  await api('/api/global/stop', { method: 'POST', body: JSON.stringify({}) });
  setTimeout(() => { refresh().catch(() => {}); }, 300);
}

async function runOnceGlobal() {
  await api('/api/global/run_once', { method: 'POST', body: JSON.stringify({}) });
  setTimeout(() => { refresh().catch(() => {}); }, 300);
}

async function cleanCache() {
  const statusEl = document.getElementById('cache-status');
  if (statusEl) statusEl.textContent = '清理中...';
  
  try {
    const result = await api('/api/clean_cache', { method: 'POST', body: JSON.stringify({}) });
    if (statusEl) statusEl.textContent = result.message || '清理完成';
    setTimeout(() => {
      if (statusEl) statusEl.textContent = '';
    }, 5000);
  } catch (err) {
    if (statusEl) statusEl.textContent = '清理失败: ' + err.message;
  }
}

refresh().catch(err => {
  const el = document.getElementById('global-error');
  if (el) el.innerHTML = `<pre>${esc(err.stack||err.message||String(err))}</pre>`;
});

loadUserContactDept().catch(() => {});
loadWenzongUserContactDept().catch(() => {});
loadWenzongConfig().catch(() => {});

try {
  const elEnabled = document.getElementById('wenzong-enabled');
  if (elEnabled) {
    elEnabled.addEventListener('change', () => {
      saveWenzongEnabledOnly();
    });
  }
} catch (e) {}

setInterval(() => {
  refresh().catch(() => {});
}, 5000);

setInterval(() => {
  const nextEl = document.getElementById('next-run');
  if (!nextEl) return;
  if (typeof window.__nextRunRemain !== 'number') return;
  if (window.__nextRunRemain <= 0) return;
  window.__nextRunRemain -= 1;
  nextEl.textContent = `距离下一次运行: ${window.__nextRunRemain} 秒`;
}, 1000);
</script>
</body>
</html>"""

    # Note: The HTML template is authored inside a Python string and previously used
    # \" sequences in many places. Browsers do not treat backslash as an escape in HTML,
    # which may break attribute parsing and prevent JS from running.
    index_html = index_html.replace('\\"', '"')

    index_html = index_html.replace("%SERVER_VERSION%", SERVER_VERSION)

    class Handler(BaseHTTPRequestHandler):
        def log_message(self, format: str, *args: Any) -> None:
            return

        def _check_auth(self) -> bool:
            if not token:
                return True
            parsed = urlparse(self.path)
            qs = parse_qs(parsed.query)
            t = (qs.get("token") or [""])[0]
            if t == token:
                return True
            hdr = (self.headers.get("X-Token") or "").strip()
            return hdr == token

        def do_GET(self) -> None:
            parsed = urlparse(self.path)
            path = parsed.path
            if path == "/favicon.ico":
                self.send_response(204)
                self.end_headers()
                return
            if not self._check_auth():
                _json_response(self, 401, {"ok": False, "error": "unauthorized"})
                return

            if path == "/":
                qs = parse_qs(parsed.query)
                t = (qs.get("token") or [""])[0]
                html = index_html
                html = html.replace("%TOKEN%", _html_escape(t))
                _html_response(self, 200, html)
                return

            if path == "/debug/html":
                _html_response(self, 200, html)
                return

            if path == "/api/global/status":
                niu_runtime = _load_niu_runtime_status(args)
                niu_runtime_items_raw = niu_runtime.get("items") if isinstance(niu_runtime, dict) else {}
                niu_runtime_items: List[Dict[str, Any]] = []
                if isinstance(niu_runtime_items_raw, dict):
                    for key in sorted(niu_runtime_items_raw.keys()):
                        item = niu_runtime_items_raw.get(key)
                        if not isinstance(item, dict):
                            continue
                        merged = {"key": key}
                        merged.update(item)
                        niu_runtime_items.append(merged)
                kuaishou_runtime = _load_kuaishou_runtime_status(args)
                kuaishou_runtime_items_raw = kuaishou_runtime.get("items") if isinstance(kuaishou_runtime, dict) else {}
                kuaishou_runtime_items: List[Dict[str, Any]] = []
                if isinstance(kuaishou_runtime_items_raw, dict):
                    for key in sorted(kuaishou_runtime_items_raw.keys()):
                        item = kuaishou_runtime_items_raw.get(key)
                        if not isinstance(item, dict):
                            continue
                        merged = {"key": key}
                        merged.update(item)
                        kuaishou_runtime_items.append(merged)
                feishu_runtime = _load_feishu_runtime_status(args)
                feishu_runtime_items_raw = feishu_runtime.get("items") if isinstance(feishu_runtime, dict) else {}
                feishu_runtime_items: List[Dict[str, Any]] = []
                if isinstance(feishu_runtime_items_raw, dict):
                    for key in sorted(feishu_runtime_items_raw.keys()):
                        item = feishu_runtime_items_raw.get(key)
                        if not isinstance(item, dict):
                            continue
                        merged = {"key": key}
                        merged.update(item)
                        feishu_runtime_items.append(merged)
                s = {
                    "running": global_runner.is_running(),
                    "last_started_at": _fmt_ts(global_runner.last_started_at),
                    "last_finished_at": _fmt_ts(global_runner.last_finished_at),
                    "last_returncode": global_runner.last_returncode,
                    "last_error": global_runner.last_error,
                    "last_log_path": global_runner.last_log_path,
                }
                next_run_in_seconds = global_runner.get_next_run_in_seconds() if global_runner.is_running() else None
                _json_response(
                    self,
                    200,
                    {
                        "ok": True,
                        "server_time": _fmt_ts(_now_epoch()),
                        "server_version": SERVER_VERSION,
                        "status": s,
                        "interval_seconds": global_runner.get_interval_seconds(),
                        "next_run_in_seconds": next_run_in_seconds,
                        "night_sleep_enabled": global_runner.night_sleep_enabled,
                        "sleeping": global_runner.is_sleeping(),
                        "accounts": accounts,
                        "niu_runtime_items": niu_runtime_items,
                        "kuaishou_runtime_items": kuaishou_runtime_items,
                        "feishu_runtime_items": feishu_runtime_items,
                    },
                )
                return

            if path == "/api/wenzong/config":
                wz = cfg.get("wenzong") or {}
                _json_response(
                    self,
                    200,
                    {
                        "ok": True,
                        "enabled": bool(wz.get("enabled", False)),
                        "classroom_url": str(wz.get("classroom_url", "") or ""),
                    },
                )
                return

            if path == "/api/copy_table":
                data = _load_copy_table_cache()
                items = data.get("items")
                if not isinstance(items, list):
                    items = []
                _json_response(self, 200, {"ok": True, "items": items, "updated_at": data.get("updated_at", "")})
                return

            if path == "/api/user_contact":
                csv_path = _user_contact_csv_path(cfg, args)
                rows = _load_user_contact_rows(csv_path)
                _json_response(self, 200, {"ok": True, "rows": rows})
                return

            if path == "/api/niu_table":
                rows = _load_niu_table(cfg, args)
                status_data = _load_niu_bg_status(cfg, args)
                _json_response(self, 200, {"ok": True, "rows": _merge_niu_rows_with_bg_status(rows, status_data)})
                return

            if path == "/api/open_niu_url":
                qs = parse_qs(parsed.query)
                idx_s = (qs.get("idx") or [""])[0].strip()
                try:
                    idx = int(idx_s)
                except Exception:
                    idx = -1
                rows = _load_niu_table(cfg, args)
                if idx < 0 or idx >= len(rows):
                    _json_response(self, 400, {"ok": False, "error": "invalid_idx"})
                    return
                url = str((rows[idx] or {}).get("url") or "").strip()
                if not url:
                    _json_response(self, 400, {"ok": False, "error": "empty_url"})
                    return

                # Use Playwright Chromium persistent profile (dedicated dir; do NOT mix with system Chrome)
                profile_dir = _niu_front_profile_dir(idx)
                os.makedirs(profile_dir, exist_ok=True)

                def _open_browser() -> None:
                    try:
                        log_dir = os.path.join(".state", "logs")
                        os.makedirs(log_dir, exist_ok=True)
                        log_path = os.path.join(log_dir, f"open_niu_slot_{idx}.log")
                        cmd = [
                            sys.executable,
                            "open_niu_browser.py",
                            "--profile-dir",
                            os.path.abspath(profile_dir),
                            "--url",
                            url,
                        ]
                        with open(log_path, "a", encoding="utf-8") as f:
                            f.write(f"\n[web] cmd={cmd!r}\n")
                            f.flush()
                            proc = subprocess.Popen(
                                cmd,
                                stdout=f,
                                stderr=f,
                                start_new_session=True,
                            )
                            f.write(f"[web] started pid={proc.pid}\n")
                            f.flush()
                        print(
                            f"[web] Started niu browser process idx={idx} pid={proc.pid} log={log_path}",
                            file=sys.stderr,
                        )
                    except Exception as e:
                        print(f"[web] Failed to start niu browser process: {e}", file=sys.stderr)

                threading.Thread(target=_open_browser, daemon=True).start()
                _json_response(
                    self,
                    200,
                    {
                        "ok": True,
                        "profile_dir": profile_dir,
                        "url": url,
                        "log_path": os.path.join(".state", "logs", f"open_niu_slot_{idx}.log"),
                    },
                )
                return

            if path == "/api/niu_login_status":
                qs = parse_qs(parsed.query)
                idx_s = (qs.get("idx") or [""])[0].strip()
                try:
                    idx = int(idx_s)
                except Exception:
                    idx = -1
                if idx < 0:
                    _json_response(self, 400, {"ok": False, "error": "invalid_idx"})
                    return
                _json_response(self, 200, {"ok": True, "status": _get_niu_status(idx)})
                return

            if path == "/api/user_contact_config":
                qs = parse_qs(parsed.query)
                scope = (qs.get("scope") or [""])[0].strip().lower()
                if scope == "wenzong":
                    dept = _load_wenzong_user_contact_dept(cfg, args)
                else:
                    dept = _load_user_contact_dept(cfg, args)
                _json_response(self, 200, {"ok": True, "dept": dept})
                return

            if path == "/api/user_contact/export_xlsx":
                csv_path = _user_contact_csv_path(cfg, args)
                rows = _load_user_contact_rows(csv_path)
                qs = parse_qs(parsed.query)
                scope = (qs.get("scope") or [""])[0].strip().lower()
                if scope == "wenzong":
                    export_dir = _wenzong_user_contact_export_dir(cfg, args)
                else:
                    export_dir = _user_contact_export_dir(cfg, args)
                info = _export_user_contact_xlsx_to_dir(cfg, args, rows, export_dir=export_dir) if export_dir else {}
                name = (info.get("name") or f"{_now_ts()}.xlsx")
                try:
                    if info.get("ok"):
                        with open(str(info.get("out_path")), "rb") as f:
                            payload = f.read()
                    else:
                        dept = _load_user_contact_dept(cfg, args)
                        rows2 = _apply_dept(rows, dept)
                        payload = _xlsx_bytes_from_rows_template(cfg, args, rows2)
                except Exception as e:
                    print(f"[web] template export failed, fallback to simple workbook: {e}", file=sys.stderr)
                    dept = _load_user_contact_dept(cfg, args)
                    rows2 = _apply_dept(rows, dept)
                    payload = _xlsx_bytes_from_rows(rows2)

                _bytes_response(
                    self,
                    200,
                    payload,
                    content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    filename=name,
                )
                return

            if path == "/api/mapping":
                qs = parse_qs(parsed.query)
                scope = (qs.get("scope") or [""])[0].strip().lower()
                if scope == "wenzong":
                    anchor_map_csv = _wenzong_anchor_map_csv_path(cfg, args)
                else:
                    anchor_map_csv = _anchor_map_csv_path(cfg, args)
                if not anchor_map_csv:
                    _json_response(self, 400, {"ok": False, "error": "no_mapping_file_configured"})
                    return
                inactive_csv = anchor_map_csv.replace(".csv", "_暂不使用.csv")

                header = ["直播账号", "快手ID", "手机号码", "密码", "主播"]
                if not os.path.exists(anchor_map_csv):
                    active_rows = []
                else:
                    active_rows = _read_csv_rows(anchor_map_csv)
                inactive_rows: List[List[str]] = []
                if os.path.exists(inactive_csv):
                    inactive_rows = _read_csv_rows(inactive_csv)

                _json_response(
                    self,
                    200,
                    {
                        "ok": True,
                        "active_rows": active_rows,
                        "inactive_rows": inactive_rows,
                        "anchor_map_csv": anchor_map_csv,
                        "inactive_csv": inactive_csv,
                    },
                )
                return

            if path == "/api/open_kuaishou_classroom":
                # 打开快手课堂（带登录状态的浏览器）
                qs = parse_qs(parsed.query)
                scope = (qs.get("scope") or [""])[0].strip().lower()
                account = (qs.get("account") or [""])[0]
                ks_id = (qs.get("ks_id") or [""])[0]
                
                if not account or not ks_id:
                    _json_response(self, 400, {"ok": False, "error": "missing account or ks_id"})
                    return
                
                # 使用快手ID作为profile目录
                profile_dir = _kuaishou_front_profile_dir(ks_id)
                classroom_url = "https://kt.kuaishou.com/student-management/offsite-student-management"
                if scope == "wenzong":
                    wz = cfg.get("wenzong") or {}
                    classroom_url = str(wz.get("classroom_url") or classroom_url)
                
                # 使用独立脚本在后台打开浏览器
                # 这样浏览器进程独立运行，不会随着API请求结束而关闭
                def _open_browser():
                    try:
                        cmd = [
                            sys.executable,
                            "open_kuaishou_browser.py",
                            "--profile-dir", os.path.abspath(profile_dir),
                            "--url", classroom_url,
                            "--account", account,
                            "--ks-id", ks_id,
                        ]
                        
                        # 使用Popen在后台运行，不等待完成
                        proc = subprocess.Popen(
                            cmd,
                            stdout=subprocess.PIPE,
                            stderr=subprocess.PIPE,
                            start_new_session=True,  # 创建新的进程组，独立于父进程
                        )
                        
                        print(f"[web] Started browser process for account={account}, ks_id={ks_id}, pid={proc.pid}", file=sys.stderr)
                        
                    except Exception as e:
                        print(f"[web] Failed to start browser process: {e}", file=sys.stderr)
                
                threading.Thread(target=_open_browser, daemon=True).start()
                
                _json_response(self, 200, {
                    "ok": True,
                    "message": f"Opening Kuaishou classroom for {account}",
                    "account": account,
                    "ks_id": ks_id,
                    "profile_dir": profile_dir
                })
                return

            if path == "/api/kuaishou_login_status":
                qs = parse_qs(parsed.query)
                scope = (qs.get("scope") or [""])[0].strip().lower()
                ks_id = (qs.get("ks_id") or [""])[0].strip()
                if not ks_id:
                    _json_response(self, 400, {"ok": False, "error": "missing_ks_id"})
                    return
                _json_response(self, 200, {"ok": True, "status": _get_kuaishou_status(scope, ks_id)})
                return

            if path == "/api/kuaishou_login_qr":
                qs = parse_qs(parsed.query)
                ks_id = (qs.get("ks_id") or [""])[0].strip()
                if not ks_id:
                    _json_response(self, 400, {"ok": False, "error": "missing_ks_id"})
                    return
                qr_path = _kuaishou_bg_qr_path(ks_id)
                qr_abs = _resolve_path(args.config, qr_path)
                if not os.path.exists(qr_abs):
                    _json_response(self, 404, {"ok": False, "error": "qr_not_found"})
                    return
                with open(qr_abs, "rb") as f:
                    payload = f.read()
                _bytes_response(self, 200, payload, content_type="image/png")
                return

            _json_response(self, 404, {"ok": False, "error": "not_found"})

        def do_POST(self) -> None:
            if not self._check_auth():
                _json_response(self, 401, {"ok": False, "error": "unauthorized"})
                return

            parsed = urlparse(self.path)
            path = parsed.path
            try:
                length = int(self.headers.get("Content-Length", "0"))
            except Exception:
                length = 0
            raw = self.rfile.read(length) if length else b""
            try:
                body = json.loads(raw.decode("utf-8")) if raw else {}
            except Exception:
                body = {}

            if path == "/api/global/start":
                global_runner.start()
                _json_response(self, 200, {"ok": True, "running": True})
                return

            if path == "/api/global/stop":
                global_runner.stop()
                _json_response(self, 200, {"ok": True, "running": False})
                return

            if path == "/api/global/run_once":
                threading.Thread(target=global_runner.run_once, daemon=True).start()
                _json_response(self, 200, {"ok": True})
                return

            if path == "/api/global/set_interval":
                interval_seconds = body.get("interval_seconds")
                if not isinstance(interval_seconds, (int, float)) or interval_seconds < 1 or interval_seconds > 3600:
                    _json_response(self, 400, {"ok": False, "error": "interval_seconds must be between 1 and 3600"})
                    return

                interval_seconds_i = int(interval_seconds)
                global_runner.set_interval(interval_seconds_i)
                
                # 保存到配置文件
                try:
                    cfg.setdefault("web", {})
                    cfg["web"]["interval_seconds"] = interval_seconds_i
                    save_json(args.config, cfg)
                except Exception as e:
                    print(f"[web] failed to save interval to config: {e}", file=sys.stderr)
                
                _json_response(self, 200, {
                    "ok": True,
                    "interval_seconds": global_runner.get_interval_seconds(),
                    "message": f"Interval set to {global_runner.get_interval_seconds()} seconds"
                })
                return

            if path == "/api/global/night_sleep":
                enabled = bool(body.get("enabled", False))
                global_runner.night_sleep_enabled = enabled
                try:
                    cfg.setdefault("web", {})
                    cfg["web"]["night_sleep_enabled"] = enabled
                    save_json(args.config, cfg)
                except Exception as e:
                    print(f"[web] failed to save night_sleep to config: {e}", file=sys.stderr)
                _json_response(self, 200, {
                    "ok": True,
                    "night_sleep_enabled": global_runner.night_sleep_enabled,
                    "sleeping": global_runner.is_sleeping(),
                })
                return

            if path == "/api/wenzong/config":
                enabled = bool(body.get("enabled", False))
                classroom_url = str(body.get("classroom_url", "") or "").strip()
                try:
                    cfg.setdefault("wenzong", {})
                    cfg["wenzong"]["enabled"] = enabled
                    if classroom_url:
                        cfg["wenzong"]["classroom_url"] = classroom_url
                    save_json(args.config, cfg)
                except Exception as e:
                    _json_response(self, 500, {"ok": False, "error": str(e)})
                    return
                _json_response(self, 200, {"ok": True, "enabled": enabled, "classroom_url": classroom_url})
                return

            if path == "/api/niu_table":
                rows = body.get("rows")
                if not isinstance(rows, list):
                    _json_response(self, 400, {"ok": False, "error": "invalid_rows"})
                    return
                try:
                    _save_niu_table(cfg, args, rows)
                except Exception as e:
                    _json_response(self, 500, {"ok": False, "error": str(e)})
                    return
                _json_response(self, 200, {"ok": True})
                return

            if path == "/api/niu_login_background":
                try:
                    idx = int(body.get("idx"))
                except Exception:
                    idx = -1
                phone = str(body.get("phone") or "").strip()
                password = str(body.get("password") or "").strip()
                rows = _load_niu_table(cfg, args)
                if idx < 0 or idx >= len(rows):
                    _json_response(self, 400, {"ok": False, "error": "invalid_idx"})
                    return
                if not phone or not password:
                    _json_response(self, 400, {"ok": False, "error": "missing_phone_or_password"})
                    return
                url = str((rows[idx] or {}).get("url") or "").strip()
                if not url:
                    _json_response(self, 400, {"ok": False, "error": "empty_url"})
                    return

                current = _get_niu_status(idx)
                if current.get("running"):
                    _json_response(self, 200, {"ok": True, "status": current})
                    return

                profile_dir = os.path.abspath(_niu_bg_profile_dir(idx))
                _set_niu_status(
                    idx,
                    {
                        "ok": False,
                        "running": True,
                        "progress": 1,
                        "message": "后台登录任务已创建",
                        "profile_dir": profile_dir,
                        "last_started_at": _now_iso(),
                    },
                    persist=True,
                )

                def _worker() -> None:
                    try:
                        def _progress(progress: int, message: str) -> None:
                            _set_niu_status(
                                idx,
                                {
                                    "running": True,
                                    "progress": int(progress),
                                    "message": str(message),
                                    "profile_dir": profile_dir,
                                },
                                persist=True,
                            )

                        result = _run_niu_background_login(
                            url=url,
                            profile_dir=profile_dir,
                            phone=phone,
                            password=password,
                            update_status=_progress,
                        )
                        _set_niu_status(
                            idx,
                            {
                                "ok": True,
                                "running": False,
                                "progress": 100,
                                "message": "后台登录完成",
                                "profile_dir": profile_dir,
                                "last_success_at": _now_iso(),
                                "finished_at": result.get("finished_at") or _now_iso(),
                                "verified": bool(result.get("verified")),
                                "url": str(result.get("url") or url),
                            },
                            persist=True,
                        )
                    except Exception as e:
                        _set_niu_status(
                            idx,
                            {
                                "ok": False,
                                "running": False,
                                "progress": 100,
                                "message": f"后台登录失败: {e}",
                                "profile_dir": profile_dir,
                                "last_error_at": _now_iso(),
                            },
                            persist=True,
                        )

                threading.Thread(target=_worker, name=f"niu-bg-login-{idx}", daemon=True).start()
                _json_response(self, 200, {"ok": True, "status": _get_niu_status(idx)})
                return

            if path == "/api/kuaishou_login_background":
                scope = str(body.get("scope") or "").strip().lower()
                account = str(body.get("account") or "").strip()
                ks_id = str(body.get("ks_id") or "").strip()
                if not account or not ks_id:
                    _json_response(self, 400, {"ok": False, "error": "missing_account_or_ks_id"})
                    return

                current = _get_kuaishou_status(scope, ks_id)
                if current.get("running"):
                    _json_response(self, 200, {"ok": True, "status": current})
                    return

                profile_dir = os.path.abspath(_kuaishou_bg_profile_dir(ks_id))
                classroom_url = "https://kt.kuaishou.com/student-management/offsite-student-management"
                if scope == "wenzong":
                    wz = cfg.get("wenzong") or {}
                    classroom_url = str(wz.get("classroom_url") or classroom_url)

                _set_kuaishou_status(
                    scope,
                    ks_id,
                    {
                        "account": account,
                        "ks_id": ks_id,
                        "scope": scope,
                        "ok": False,
                        "running": True,
                        "progress": 1,
                        "message": "后台登录任务已创建",
                        "profile_dir": profile_dir,
                        "last_started_at": _now_iso(),
                    },
                    persist=True,
                )

                def _worker() -> None:
                    try:
                        def _progress(payload: Dict[str, Any]) -> None:
                            patch = {
                                "account": account,
                                "ks_id": ks_id,
                                "scope": scope,
                                "profile_dir": profile_dir,
                                "running": True,
                            }
                            if isinstance(payload, dict):
                                patch.update(payload)
                            _set_kuaishou_status(scope, ks_id, patch, persist=True)

                        result = _run_kuaishou_background_login(
                            profile_dir=profile_dir,
                            classroom_url=classroom_url,
                            ks_id=ks_id,
                            update_status=_progress,
                        )
                        _set_kuaishou_status(
                            scope,
                            ks_id,
                            {
                                "account": account,
                                "ks_id": ks_id,
                                "scope": scope,
                                "ok": True,
                                "running": False,
                                "progress": 100,
                                "message": "后台登录完成",
                                "profile_dir": profile_dir,
                                "finished_at": result.get("finished_at") or _now_iso(),
                                "url": str(result.get("url") or classroom_url),
                                "qr_path": str(result.get("qr_path") or ""),
                            },
                            persist=True,
                        )
                    except Exception as e:
                        _set_kuaishou_status(
                            scope,
                            ks_id,
                            {
                                "account": account,
                                "ks_id": ks_id,
                                "scope": scope,
                                "ok": False,
                                "running": False,
                                "progress": 100,
                                "message": f"后台登录失败: {e}",
                                "profile_dir": profile_dir,
                                "last_error_at": _now_iso(),
                            },
                            persist=True,
                        )

                threading.Thread(target=_worker, name=f"kuaishou-bg-login-{ks_id}", daemon=True).start()
                _json_response(self, 200, {"ok": True, "status": _get_kuaishou_status(scope, ks_id)})
                return

            if path == "/api/mapping":
                # 保存主播映射表（两个文件）
                active_rows = body.get("active_rows")
                inactive_rows = body.get("inactive_rows")
                
                if not isinstance(active_rows, list) or not isinstance(inactive_rows, list):
                    _json_response(self, 400, {"ok": False, "error": "invalid_rows"})
                    return

                parsed = urlparse(self.path)
                qs = parse_qs(parsed.query)
                scope = (qs.get("scope") or [""])[0].strip().lower()

                if scope == "wenzong":
                    anchor_map_csv = _wenzong_anchor_map_csv_path(cfg, args)
                else:
                    anchor_map_csv_raw = (cfg.get("mapping") or {}).get("anchor_map_csv") or ""
                    anchor_map_csv = _resolve_path(args.config, anchor_map_csv_raw)
                if not anchor_map_csv:
                    _json_response(self, 400, {"ok": False, "error": "no_mapping_file_configured"})
                    return

                try:
                    # 表头
                    header = ["直播账号", "快手ID", "手机号码", "密码", "主播"]
                    
                    # 保存活跃表
                    with open(anchor_map_csv, "w", encoding="utf-8-sig", newline="") as f:
                        writer = csv.writer(f)
                        writer.writerow(header)  # 写入表头
                        for row in active_rows:
                            if isinstance(row, list) and row != header:  # 跳过表头行
                                writer.writerow(row)
                    
                    # 保存暂不使用表
                    inactive_csv = anchor_map_csv.replace(".csv", "_暂不使用.csv")
                    with open(inactive_csv, "w", encoding="utf-8-sig", newline="") as f:
                        writer = csv.writer(f)
                        writer.writerow(header)  # 写入表头
                        for row in inactive_rows:
                            if isinstance(row, list) and row != header:  # 跳过表头行
                                writer.writerow(row)
                    
                    _json_response(self, 200, {"ok": True, "message": "saved"})
                except Exception as e:
                    _json_response(self, 500, {"ok": False, "error": str(e)})
                return

            if path == "/api/user_contact":
                rows = body.get("rows")
                if not isinstance(rows, list):
                    _json_response(self, 400, {"ok": False, "error": "invalid_rows"})
                    return
                csv_path = _user_contact_csv_path(cfg, args)
                try:
                    _save_user_contact_rows(csv_path, rows)
                    export_dir = _user_contact_export_dir(cfg, args)
                    export_info = {}
                    if export_dir:
                        export_info = _export_user_contact_xlsx_to_dir(cfg, args, rows, export_dir=export_dir)
                    _json_response(
                        self,
                        200,
                        {"ok": True, "message": "saved", "export": export_info},
                    )
                except Exception as e:
                    _json_response(self, 500, {"ok": False, "error": str(e)})
                return

            if path == "/api/user_contact_config":
                dept = (body.get("dept") or "").strip()
                try:
                    parsed = urlparse(self.path)
                    qs = parse_qs(parsed.query)
                    scope = (qs.get("scope") or [""])[0].strip().lower()
                    if scope == "wenzong":
                        _save_wenzong_user_contact_dept(cfg, args, dept)
                    else:
                        _save_user_contact_dept(cfg, args, dept)
                    _json_response(self, 200, {"ok": True, "dept": dept})
                except Exception as e:
                    _json_response(self, 500, {"ok": False, "error": str(e)})
                return

            if path == "/api/user_contact/import_xlsx":
                try:
                    _, file_bytes = _read_multipart_file(raw, self.headers.get("Content-Type", ""))
                    rows = _parse_xlsx_rows(file_bytes)
                    csv_path = _user_contact_csv_path(cfg, args)
                    _save_user_contact_rows(csv_path, rows)
                    export_dir = _user_contact_export_dir(cfg, args)
                    export_info = {}
                    if export_dir:
                        export_info = _export_user_contact_xlsx_to_dir(cfg, args, rows, export_dir=export_dir)
                    _json_response(
                        self,
                        200,
                        {"ok": True, "message": "imported", "count": len(rows), "export": export_info},
                    )
                except Exception as e:
                    _json_response(self, 400, {"ok": False, "error": str(e)})
                return

            if path == "/api/copy_table/mark_copied":
                key = (body.get("key") or "").strip()
                copied = bool(body.get("copied", True))
                ok = _set_copy_table_item_copied(key, copied)
                if not ok:
                    _json_response(self, 400, {"ok": False, "error": "not_found"})
                    return
                _json_response(self, 200, {"ok": True})
                return

            if path == "/api/copy_table/clear":
                _clear_copy_table_cache()
                _json_response(self, 200, {"ok": True})
                return

            if path == "/api/clean_cache":
                # 清理浏览器缓存
                try:
                    profile_dirs = [
                        '.state/kuaishou_profiles',
                        '.state/kuaishou_profiles_front',
                        '.state/kuaishou_profiles_bg',
                        '.state/niu_chrome_profile', 
                        '.state/feishu_profile',
                        '.state/niu_profile',
                        '.state/kuaishou_chromium',
                        '.state/niu_pw_profiles',
                        '.state/niu_pw_profiles_front',
                        '.state/niu_pw_profiles_bg',
                    ]

                    total_freed = _clean_browser_profile_caches(profile_dirs)

                    # Also prune old web logs to keep disk stable.
                    try:
                        keep_web_logs = 10
                        try:
                            keep_web_logs = int((cfg.get('mapping') or {}).get('keep_web_logs', 10) or 10)
                        except Exception:
                            keep_web_logs = 10
                        if keep_web_logs < 1:
                            keep_web_logs = 1
                        _cleanup_keep_latest_files(os.path.join('.state', 'logs', 'web'), keep=keep_web_logs, exts=['.log'])
                    except Exception:
                        pass

                    freed_gb = float(total_freed) / 1024 / 1024 / 1024
                    _json_response(self, 200, {
                        "ok": True, 
                        "message": f"清理完成，释放 {freed_gb:.2f}GB 空间",
                        "freed_bytes": total_freed
                    })
                except Exception as e:
                    _json_response(self, 500, {"ok": False, "error": str(e)})
                return

            _json_response(self, 404, {"ok": False, "error": "not_found"})

    return Handler


def main() -> None:
    p = argparse.ArgumentParser()
    p.add_argument("--config", default="config.json")
    p.add_argument("--host", default="0.0.0.0")
    p.add_argument("--port", type=int, default=8010)
    p.add_argument("--interval-seconds", type=int, default=120)
    p.add_argument("--headless", action="store_true")
    p.add_argument("--token", default="")
    p.add_argument("--keep-exports", type=int, default=20)
    p.add_argument("--keep-sync-logs", type=int, default=50)
    p.add_argument("--keep-web-logs", type=int, default=50)
    args = p.parse_args()

    repo_dir = os.path.dirname(os.path.abspath(__file__))
    config_path = args.config
    if not os.path.isabs(config_path):
        config_path = os.path.join(repo_dir, config_path)

    cfg = load_json(config_path)
    anchor_map_csv_raw = (cfg.get("mapping") or {}).get("anchor_map_csv") or ""
    anchor_map_csv = _resolve_path(config_path, anchor_map_csv_raw)
    if not anchor_map_csv:
        raise RuntimeError(f"Invalid anchor_map_csv: {anchor_map_csv_raw}")

    web_cfg = cfg.get("web") or {}
    rotate_on_start = bool(web_cfg.get("rotate_token_on_start", False))

    token = (args.token or web_cfg.get("token") or os.environ.get("OPENCLAW_WEB_TOKEN") or "").strip()
    if rotate_on_start:
        token = secrets.token_urlsafe(18)
        cfg.setdefault("web", {})
        cfg["web"]["token"] = token
        save_json(config_path, cfg)
        print(f"[web] rotated token and saved to config: token={token}", file=sys.stderr)

    # 从配置文件读取间隔时间，如果没有则使用命令行参数或默认值
    interval_seconds_cfg = web_cfg.get("interval_seconds")
    interval_minutes_legacy = web_cfg.get("interval_minutes")
    if interval_seconds_cfg is not None:
        interval_seconds = int(interval_seconds_cfg)
        print(f"[web] using interval from config: {interval_seconds} seconds", file=sys.stderr)
    elif interval_minutes_legacy is not None:
        interval_seconds = int(interval_minutes_legacy) * 60
        print(f"[web] using interval from legacy config: {interval_seconds} seconds", file=sys.stderr)
    else:
        interval_seconds = int(args.interval_seconds)
        print(f"[web] using default interval: {interval_seconds} seconds", file=sys.stderr)

    state_dir = os.path.join(repo_dir, ".state")
    log_dir = os.path.join(state_dir, "logs", "web")
    global_run_lock = threading.Lock()

    exports_root = os.path.join(repo_dir, "exports")
    sync_logs_dir = os.path.join(repo_dir, "logs")

    effective_headless = bool(args.headless or bool(web_cfg.get("headless", False)))

    global_runner = GlobalRunner(
        repo_dir=repo_dir,
        config_path=config_path,
        interval_seconds=interval_seconds,
        headless=effective_headless,
        global_run_lock=global_run_lock,
        log_dir=log_dir,
        sync_logs_dir=sync_logs_dir,
        keep_sync_logs=args.keep_sync_logs,
        keep_web_logs=args.keep_web_logs,
    )
    global_runner.night_sleep_enabled = bool(web_cfg.get("night_sleep_enabled", False))
    print(f"[web] night_sleep_enabled={global_runner.night_sleep_enabled}", file=sys.stderr)
    print(f"[web] kuaishou_headless={effective_headless}", file=sys.stderr)

    handler_cls = build_handler(
        token=token, 
        accounts=[],
        runners={},
        global_runner=global_runner,
        cfg=cfg, 
        args=args,
        reload_callback=lambda: None,
    )
    httpd = HTTPServer((args.host, args.port), handler_cls)

    print(f"[web] listening on http://{args.host}:{args.port}", file=sys.stderr)
    if token:
        print(f"[web] open: http://127.0.0.1:{args.port}/?token={token}", file=sys.stderr)
    else:
        print(f"[web] auth disabled (no token). open: http://127.0.0.1:{args.port}/", file=sys.stderr)

    try:
        httpd.serve_forever()
    except KeyboardInterrupt:
        pass
    finally:
        global_runner.stop()
        httpd.server_close()


if __name__ == "__main__":
    main()
